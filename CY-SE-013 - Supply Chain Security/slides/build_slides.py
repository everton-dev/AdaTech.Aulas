# -*- coding: utf-8 -*-
"""
build_slides.py — CY-SE-013 · Supply Chain Security
Gera 9 apresentações (8 aulas + encerramento) no layout Ada usado em
arquitetura-de-servicos/slides: fundo cream, barra gradiente, cards de
cronograma com o dia atual em destaque, bullets numerados, diagramas nativos,
código em janela de terminal e recap. Toda aula tem notas do apresentador.

Cronograma: Seg/Ter/Sex a partir de Sex 16/10/2026. 02/11 (Finados) sem aula.
Último encontro: Sex 06/11/2026 — Concluir atividades e preencher rubricas.
"""
import os
import unicodedata
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(BASE, "Ada_Template PPT.pptx")
IMG = os.path.join(BASE, "imagens")
ADA_LOGO = os.path.join(IMG, "brand_ada_logo.png")
GRAD_H = os.path.join(IMG, "brand_gradient_h.png")
SWIRL = os.path.join(IMG, "brand_swirl_green.png")

C_DARK = RGBColor(0x18, 0x1C, 0x23)
C_BLUE = RGBColor(0x23, 0x71, 0xFF)
C_GREEN = RGBColor(0xA6, 0xF7, 0x50)
C_GREENL = RGBColor(0xDA, 0xF9, 0xDA)
C_RED = RGBColor(0xFF, 0x23, 0x4F)
C_WHITE = RGBColor(0xFF, 0xFD, 0xF8)
C_GRAY = RGBColor(0x6A, 0x70, 0x7D)
C_LGRAY = RGBColor(0xEC, 0xED, 0xF3)
C_TERM = RGBColor(0x2A, 0x2F, 0x3A)

F_LIGHT = "Segoe UI Light"
F_REG = "Segoe UI"
F_MED = "Segoe UI Semibold"
F_CODE = "Consolas"

SW, SH = 9144000, 5143500
MX = 311700                      # margem lateral
FOOT_H = 360000                  # área reservada ao rodapé

MODULE = "CY-SE-013 - Supply Chain Security"
PROFESSOR = "Everton J Benedicto"

CRONOGRAMA = [
    ("01", "Sex 16/10", "Policy as Code e Compliance Automatizado"),
    ("02", "Seg 19/10", "Linguagem Rego e Escrita de Políticas"),
    ("03", "Ter 20/10", "Enforcement: OPA Gatekeeper e Admission Control"),
    ("04", "Sex 23/10", "SBOM: Fundamentos, SPDX e CycloneDX"),
    ("05", "Seg 26/10", "Análise de SBOM e Gestão de Dependências"),
    ("06", "Ter 27/10", "Supply Chain Security: Ataques e Frameworks"),
    ("07", "Sex 30/10", "Integridade e Procedência de Artefatos"),
    ("08", "Ter 03/11", "Compliance Contínuo, Auditoria e Projeto Final"),
    ("09", "Sex 06/11", "Concluir atividades e preencher rubricas"),
]

COMBINADOS = [
    "Câmeras ligadas.",
    "Entrem no horário: início e intervalo.",
    "Participem da aula.",
    "Tirem dúvidas.",
    "Levantem a mão para tirar dúvidas.",
]


# ─── helpers ──────────────────────────────────────────────────────────────────
def _remove_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        try:
            prs.part.drop_rel(sid.get(qn("r:id")))
        except Exception:
            pass
        lst.remove(sid)


def new_prs():
    prs = Presentation(TEMPLATE)
    _remove_all_slides(prs)
    return prs


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[3])  # CUSTOM vazio, fundo cream


def _fill(s, color):
    if color is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = color


def _line(s, color, w):
    if color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = color
        s.line.width = Emu(w)


def rect(sl, x, y, w, h, fill=None, line=None, lw=12700):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    _fill(s, fill); _line(s, line, lw); s.shadow.inherit = False
    return s


def rrect(sl, x, y, w, h, fill=None, line=None, radius=0.08, lw=12700):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    _fill(s, fill); _line(s, line, lw); s.shadow.inherit = False
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    return s


def oval(sl, x, y, w, h, fill=None, line=None, lw=12700):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(y), Emu(w), Emu(h))
    _fill(s, fill); _line(s, line, lw); s.shadow.inherit = False
    return s


def arrow(sl, x, y, w, h, fill):
    s = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(x), Emu(y), Emu(w), Emu(h))
    _fill(s, fill); _line(s, None, 0); s.shadow.inherit = False
    return s


def tb(sl, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       line_spacing=None, space_after=None):
    box = sl.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after is not None:
            p.space_after = Pt(space_after)
        for text, st in para:
            r = p.add_run()
            r.text = text
            r.font.name = st.get("font", F_REG)
            r.font.size = Pt(st.get("sz", 14))
            r.font.bold = st.get("bold", False)
            r.font.italic = st.get("italic", False)
            r.font.color.rgb = st.get("color", C_DARK)
    return box


def pic(sl, path, x, y, w, h):
    return sl.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))


def pic_fit(sl, path, bx, by, bw, bh, frame=True):
    iw, ih = Image.open(path).size
    ar = iw / ih
    if ar > bw / bh:
        w = bw; h = int(bw / ar)
    else:
        h = bh; w = int(bh * ar)
    x = bx + (bw - w) // 2
    y = by + (bh - h) // 2
    if frame:
        pad = 60000
        rrect(sl, x - pad, y - pad, w + 2 * pad, h + 2 * pad, fill=C_WHITE, line=C_LGRAY, radius=0.04, lw=9525)
    return sl.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))


def set_notes(sl, text, minutes):
    sl.notes_slide.notes_text_frame.text = f"Tempo sugerido: {minutes} min\n\n{text}"


# ─── marca ────────────────────────────────────────────────────────────────────
def topbar(sl):
    iw, ih = Image.open(GRAD_H).size
    h = int(SW / (iw / ih))
    pic(sl, GRAD_H, 0, 0, SW, h)
    return h


def footer(sl, aula_titulo, idx):
    y = SH - 300000
    rect(sl, MX, y, SW - 2 * MX, 6000, fill=C_LGRAY)
    tb(sl, MX, y + 40000, SW - 2 * MX - 600000, 220000,
       [[(f"Ada · {MODULE} · {aula_titulo} | {PROFESSOR}", {"font": F_REG, "sz": 8, "color": C_GRAY})]],
       anchor=MSO_ANCHOR.MIDDLE)
    tb(sl, SW - MX - 500000, y + 40000, 500000, 220000,
       [[(str(idx), {"font": F_MED, "sz": 8, "color": C_BLUE})]],
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def header(sl, titulo, kicker=None):
    top = topbar(sl)
    y = top + 180000
    if kicker:
        tb(sl, MX, y, SW - 2 * MX, 240000, [[(kicker.upper(), {"font": F_MED, "sz": 11, "color": C_BLUE})]])
        y += 260000
    tb(sl, MX, y, SW - 2 * MX, 560000, [[(titulo, {"font": F_MED, "sz": 26, "color": C_DARK})]])
    rect(sl, MX, y + 560000, 900000, 55000, fill=C_GREEN)
    return y + 690000


def swirl(sl, w, x, y):
    iw, ih = Image.open(SWIRL).size
    pic(sl, SWIRL, x, y, w, int(w / (iw / ih)))


# ─── slides ───────────────────────────────────────────────────────────────────
def build_capa(prs, num, titulo, subtitulo, data):
    sl = add_slide(prs)
    topbar(sl)
    swirl(sl, 3400000, SW - 2700000, SH - 2400000)
    tb(sl, 560000, 560000, 2400000, 1500000,
       [[(f"{num:02d}", {"font": F_LIGHT, "sz": 120, "color": C_GREEN, "bold": True})]])
    tb(sl, 620000, 2380000, 6000000, 300000,
       [[(f"AULA {num:02d}  ·  {data}", {"font": F_MED, "sz": 13, "color": C_BLUE})]])
    tb(sl, 620000, 2660000, 6600000, 1100000,
       [[(titulo, {"font": F_MED, "sz": 34, "color": C_DARK})]], line_spacing=1.0)
    tb(sl, 620000, 3720000, 6600000, 500000,
       [[(subtitulo, {"font": F_LIGHT, "sz": 16, "color": C_GRAY, "italic": True})]])
    iw, ih = Image.open(ADA_LOGO).size
    lw = 1200000; lh = int(lw / (iw / ih))
    pic(sl, ADA_LOGO, SW - lw - 360000, SH - lh - 360000, lw, lh)
    tag_w = 4600000
    rrect(sl, 620000, 4350000, tag_w, 300000, fill=C_DARK, radius=0.5)
    tb(sl, 620000, 4350000, tag_w, 300000,
       [[("Módulo · " + MODULE, {"font": F_REG, "sz": 10, "color": C_GREEN})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tb(sl, 620000, 4720000, tag_w, 240000,
       [[("Prof.: " + PROFESSOR, {"font": F_MED, "sz": 11, "color": C_DARK})]], anchor=MSO_ANCHOR.MIDDLE)
    return sl


def build_cronograma(prs, num, titulo, idx):
    sl = add_slide(prs)
    top = topbar(sl)
    tb(sl, MX, top + 150000, SW - 2 * MX, 520000, [[("Cronograma das Aulas", {"font": F_MED, "sz": 26, "color": C_DARK})]])
    rect(sl, MX, top + 660000, 900000, 55000, fill=C_GREEN)
    grid_y = top + 850000
    cols, rows, gap = 3, 3, 150000
    card_w = (SW - 2 * MX - (cols - 1) * gap) // cols
    card_h = (SH - grid_y - FOOT_H - (rows - 1) * gap) // rows
    for i, (n, dia, t) in enumerate(CRONOGRAMA):
        x = MX + (i % cols) * (card_w + gap)
        y = grid_y + (i // cols) * (card_h + gap)
        current = (i + 1) == num
        if current:
            rrect(sl, x, y, card_w, card_h, fill=C_DARK, radius=0.06)
            rrect(sl, x, y, 70000, card_h, fill=C_GREEN, radius=0.5)
            dia_c, tit_c = C_GREEN, C_WHITE
        else:
            rrect(sl, x, y, card_w, card_h, fill=C_WHITE, line=C_LGRAY, radius=0.06, lw=9525)
            dia_c, tit_c = C_BLUE, C_DARK
        pad = 150000
        tb(sl, x + pad, y + pad - 60000, card_w - 2 * pad, 320000, [[(n, {"font": F_MED, "sz": 22, "color": C_GREEN, "bold": True})]])
        tb(sl, x + pad, y + pad + 240000, card_w - 2 * pad, 200000, [[(dia, {"font": F_MED, "sz": 10, "color": dia_c})]])
        tb(sl, x + pad, y + pad + 440000, card_w - 2 * pad, card_h - pad - 480000,
           [[(t, {"font": F_REG, "sz": 10, "color": tit_c})]], line_spacing=1.0)
    footer(sl, titulo, idx)
    return sl


def build_combinado(prs, titulo, idx):
    sl = add_slide(prs)
    top = topbar(sl)
    swirl(sl, 3000000, SW - 2500000, top + 400000)
    tb(sl, MX, top + 150000, SW - 2 * MX, 260000, [[("COMBINADO DA TURMA", {"font": F_MED, "sz": 13, "color": C_GREEN})]])
    tb(sl, MX, top + 400000, SW - 2 * MX, 700000, [[("Por favor!", {"font": F_MED, "sz": 46, "color": C_DARK})]])
    body_y = top + 1250000
    col_gap = 260000
    left_w = int((SW - 2 * MX - col_gap) * 0.38)
    right_w = SW - 2 * MX - col_gap - left_w
    right_x = MX + left_w + col_gap
    body_h = SH - body_y - FOOT_H
    rrect(sl, MX, body_y, left_w, 300000, fill=C_BLUE, radius=0.12)
    tb(sl, MX, body_y, left_w, 300000, [[("Bora participar!", {"font": F_MED, "sz": 13, "color": C_WHITE})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card_y = body_y + 380000
    card_h = body_h - 380000
    rrect(sl, MX, card_y, left_w, card_h, fill=C_WHITE, line=C_LGRAY, radius=0.05, lw=9525)
    tb(sl, MX + 150000, card_y + 200000, left_w - 300000, 900000,
       [[("Câmera ligada", {"font": F_MED, "sz": 15})], [("+ mão levantada", {"font": F_MED, "sz": 15})],
        [("= aula top!", {"font": F_MED, "sz": 15, "color": C_GREEN})]], line_spacing=1.15)
    btn_h = 360000
    btn_y = card_y + card_h - btn_h - 150000
    rrect(sl, MX + 150000, btn_y, left_w - 300000, btn_h, fill=C_GREEN, radius=0.5)
    tb(sl, MX + 150000, btn_y, left_w - 300000, btn_h, [[("INTERVALO · 15 min", {"font": F_MED, "sz": 12})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rrect(sl, right_x, body_y, right_w, 300000, fill=C_BLUE, radius=0.10)
    tb(sl, right_x, body_y, right_w, 300000, [[("Combinados da nossa aula", {"font": F_MED, "sz": 13, "color": C_WHITE})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    list_y = body_y + 400000
    each = (body_h - 400000) // len(COMBINADOS)
    for i, item in enumerate(COMBINADOS):
        iy = list_y + i * each
        cd = 240000
        oval(sl, right_x, iy + (each - cd) // 2, cd, cd, fill=C_GREEN)
        tb(sl, right_x, iy + (each - cd) // 2, cd, cd, [[(str(i + 1), {"font": F_MED, "sz": 13, "bold": True})]],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(sl, right_x + cd + 180000, iy, right_w - cd - 180000, each, [[(item, {"font": F_REG, "sz": 15})]],
           anchor=MSO_ANCHOR.MIDDLE)
    footer(sl, titulo, idx)
    return sl


def _bullet_runs(b, sz):
    if isinstance(b, tuple):
        return [[(b[0] + " ", {"font": F_MED, "sz": sz}), (b[1], {"font": F_REG, "sz": sz, "color": C_GRAY})]]
    return [[(b, {"font": F_REG, "sz": sz})]]


def _analogia(sl, ay, ah, texto, label="💡 ANALOGIA"):
    rrect(sl, MX, ay, SW - 2 * MX, ah, fill=C_GREENL, radius=0.06)
    rrect(sl, MX, ay, 70000, ah, fill=C_GREEN, radius=0.5)
    tb(sl, MX + 200000, ay + 90000, SW - 2 * MX - 400000, 240000, [[(label, {"font": F_MED, "sz": 11, "color": C_BLUE})]])
    tb(sl, MX + 200000, ay + 340000, SW - 2 * MX - 400000, ah - 400000,
       [[(texto, {"font": F_REG, "sz": 13, "italic": True})]], line_spacing=1.1)


def build_content(prs, titulo_aula, idx, titulo, bullets, analogia=None, kicker=None):
    sl = add_slide(prs)
    y = header(sl, titulo, kicker)
    body_h = SH - y - FOOT_H
    ana_h = 900000 if analogia else 0
    bullets_h = body_h - ana_h - (150000 if analogia else 0)
    each = min(bullets_h // max(len(bullets), 1), 620000)
    for i, b in enumerate(bullets):
        by = y + i * each
        cd = 300000
        oval(sl, MX, by + 30000, cd, cd, fill=C_BLUE)
        tb(sl, MX, by + 30000, cd, cd, [[(str(i + 1), {"font": F_MED, "sz": 13, "color": C_WHITE, "bold": True})]],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(sl, MX + cd + 200000, by, SW - 2 * MX - cd - 200000, each, _bullet_runs(b, 14),
           anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    if analogia:
        _analogia(sl, SH - FOOT_H - ana_h, ana_h, analogia)
    footer(sl, titulo_aula, idx)
    return sl


def build_image(prs, titulo_aula, idx, titulo, image, bullets=None, caption=None, kicker=None, fonte=None):
    sl = add_slide(prs)
    y = header(sl, titulo, kicker)
    body_h = SH - y - FOOT_H
    path = os.path.join(IMG, image)
    if bullets:
        col_gap = 300000
        img_w = int((SW - 2 * MX - col_gap) * 0.52)
        txt_w = SW - 2 * MX - col_gap - img_w
        pic_fit(sl, path, MX, y, img_w, body_h - (260000 if fonte else 0))
        if fonte:
            tb(sl, MX, y + body_h - 230000, img_w, 200000,
               [[("Fonte: " + fonte, {"font": F_REG, "sz": 7, "color": C_GRAY, "italic": True})]])
        tx = MX + img_w + col_gap
        each = min(body_h // max(len(bullets), 1), 640000)
        for i, b in enumerate(bullets):
            by = y + i * each
            cd = 260000
            oval(sl, tx, by + 20000, cd, cd, fill=C_GREEN)
            tb(sl, tx, by + 20000, cd, cd, [[(str(i + 1), {"font": F_MED, "sz": 12, "bold": True})]],
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            tb(sl, tx + cd + 160000, by, txt_w - cd - 160000, each, _bullet_runs(b, 13),
               anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    else:
        cap_h = 300000 if caption else 0
        fon_h = 230000 if fonte else 0
        pic_fit(sl, path, MX, y, SW - 2 * MX, body_h - cap_h - fon_h)
        if caption:
            tb(sl, MX, y + body_h - cap_h - fon_h, SW - 2 * MX, cap_h,
               [[(caption, {"font": F_REG, "sz": 12, "color": C_GRAY, "italic": True})]], align=PP_ALIGN.CENTER)
        if fonte:
            tb(sl, MX, SH - 460000, SW - 2 * MX, 180000,
               [[("Fonte: " + fonte, {"font": F_REG, "sz": 7, "color": C_GRAY, "italic": True})]], align=PP_ALIGN.CENTER)
    footer(sl, titulo_aula, idx)
    return sl


def build_flow(prs, titulo_aula, idx, titulo, steps, kicker=None, analogia=None):
    """Diagrama nativo: caixas em sequência com setas; steps = [(rótulo, descrição)]."""
    sl = add_slide(prs)
    y = header(sl, titulo, kicker)
    body_h = SH - y - FOOT_H
    ana_h = 900000 if analogia else 0
    n = len(steps)
    arrow_w = 260000
    gap = 90000
    box_w = (SW - 2 * MX - (n - 1) * (arrow_w + 2 * gap)) // n
    box_h = 760000
    top = y + 180000
    palette = [C_BLUE, C_DARK, C_BLUE, C_DARK, C_BLUE]
    for i, (label, desc) in enumerate(steps):
        x = MX + i * (box_w + arrow_w + 2 * gap)
        fill = palette[i % len(palette)]
        rrect(sl, x, top, box_w, box_h, fill=fill, radius=0.12)
        oval(sl, x + 110000, top + 110000, 260000, 260000, fill=C_GREEN)
        tb(sl, x + 110000, top + 110000, 260000, 260000, [[(str(i + 1), {"font": F_MED, "sz": 11, "bold": True})]],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(sl, x + 110000, top + 400000, box_w - 220000, 300000,
           [[(label, {"font": F_MED, "sz": 13 if n <= 4 else 11, "color": C_WHITE})]], anchor=MSO_ANCHOR.MIDDLE)
        tb(sl, x, top + box_h + 150000, box_w, body_h - box_h - 330000 - ana_h,
           [[(desc, {"font": F_REG, "sz": 11, "color": C_GRAY})]], line_spacing=1.05)
        if i < n - 1:
            arrow(sl, x + box_w + gap, top + box_h // 2 - 110000, arrow_w, 220000, C_GREEN)
    if analogia:
        _analogia(sl, SH - FOOT_H - ana_h, ana_h, analogia)
    footer(sl, titulo_aula, idx)
    return sl


def build_code(prs, titulo_aula, idx, titulo, intro, lines, kicker=None, lang="bash"):
    sl = add_slide(prs)
    y = header(sl, titulo, kicker)
    body_h = SH - y - FOOT_H
    if intro:
        tb(sl, MX, y, SW - 2 * MX, 320000, [[(intro, {"font": F_REG, "sz": 13, "color": C_GRAY})]], line_spacing=1.05)
        y += 380000
        body_h -= 380000
    rrect(sl, MX, y, SW - 2 * MX, body_h, fill=C_DARK, radius=0.03)
    bar_h = 260000
    rrect(sl, MX, y, SW - 2 * MX, bar_h, fill=C_TERM, radius=0.03)
    for i, col in enumerate([C_RED, C_GREEN, C_BLUE]):
        oval(sl, MX + 150000 + i * 200000, y + 80000, 100000, 100000, fill=col)
    tb(sl, MX, y, SW - 2 * MX, bar_h, [[(lang, {"font": F_CODE, "sz": 9, "color": C_GRAY})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    sz = 11 if len(lines) <= 12 else 10 if len(lines) <= 15 else 9
    paras = []
    for ln in lines:
        s = ln.strip()
        if s.startswith("#") or s.startswith("//"):
            paras.append([(ln, {"font": F_CODE, "sz": sz, "color": C_GRAY})])
        elif s == "":
            paras.append([(" ", {"font": F_CODE, "sz": 6, "color": C_DARK})])
        elif s.startswith(">"):
            paras.append([(ln, {"font": F_CODE, "sz": sz, "color": C_WHITE})])
        else:
            paras.append([(ln, {"font": F_CODE, "sz": sz, "color": C_GREEN})])
    tb(sl, MX + 200000, y + bar_h + 100000, SW - 2 * MX - 400000, body_h - bar_h - 200000, paras, line_spacing=1.12)
    footer(sl, titulo_aula, idx)
    return sl


def build_comparison(prs, titulo_aula, idx, titulo, a_title, a_items, b_title, b_items, kicker=None):
    sl = add_slide(prs)
    y = header(sl, titulo, kicker)
    body_h = SH - y - FOOT_H
    gap = 260000
    col_w = (SW - 2 * MX - gap) // 2
    for x, ct, items, accent in [(MX, a_title, a_items, C_BLUE), (MX + col_w + gap, b_title, b_items, C_GREEN)]:
        rrect(sl, x, y, col_w, body_h, fill=C_WHITE, line=C_LGRAY, radius=0.04, lw=9525)
        rrect(sl, x, y, col_w, 320000, fill=accent, radius=0.04)
        tb(sl, x + 150000, y, col_w - 300000, 320000,
           [[(ct, {"font": F_MED, "sz": 14, "color": C_WHITE if accent == C_BLUE else C_DARK})]], anchor=MSO_ANCHOR.MIDDLE)
        iy = y + 460000
        each = (body_h - 560000) // max(len(items), 1)
        for it in items:
            dot = 140000
            oval(sl, x + 150000, iy + (each - dot) // 2, dot, dot, fill=accent)
            tb(sl, x + 150000 + dot + 130000, iy, col_w - 300000 - dot - 130000, each,
               [[(it, {"font": F_REG, "sz": 12})]], anchor=MSO_ANCHOR.MIDDLE)
            iy += each
    footer(sl, titulo_aula, idx)
    return sl


def build_lab(prs, titulo_aula, idx, titulo, passos, entregavel, kicker="Mão na massa"):
    sl = add_slide(prs)
    y = header(sl, titulo, kicker)
    body_h = SH - y - FOOT_H
    ent_h = 620000
    list_h = body_h - ent_h - 150000
    each = min(list_h // max(len(passos), 1), 560000)
    for i, p in enumerate(passos):
        by = y + i * each
        h = each - 90000
        rrect(sl, MX, by, SW - 2 * MX, h, fill=C_LGRAY, radius=0.08)
        cd = 280000
        oval(sl, MX + 150000, by + (h - cd) // 2, cd, cd, fill=C_BLUE)
        tb(sl, MX + 150000, by + (h - cd) // 2, cd, cd, [[(str(i + 1), {"font": F_MED, "sz": 12, "color": C_WHITE, "bold": True})]],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(sl, MX + 150000 + cd + 160000, by, SW - 2 * MX - cd - 460000, h, _bullet_runs(p, 13),
           anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
    ay = SH - FOOT_H - ent_h
    rrect(sl, MX, ay, SW - 2 * MX, ent_h, fill=C_GREENL, radius=0.06)
    rrect(sl, MX, ay, 70000, ent_h, fill=C_GREEN, radius=0.5)
    tb(sl, MX + 200000, ay + 70000, SW - 2 * MX - 400000, 220000, [[("✅ ENTREGÁVEL", {"font": F_MED, "sz": 11, "color": C_BLUE})]])
    tb(sl, MX + 200000, ay + 290000, SW - 2 * MX - 400000, ent_h - 320000,
       [[(entregavel, {"font": F_REG, "sz": 12, "italic": True})]], line_spacing=1.05)
    footer(sl, titulo_aula, idx)
    return sl


def build_recap(prs, titulo_aula, idx, pontos):
    sl = add_slide(prs)
    y = header(sl, "Recapitulando a Aula", kicker="Resumo")
    body_h = SH - y - FOOT_H
    each = min(body_h // max(len(pontos), 1), 700000)
    for i, p in enumerate(pontos):
        by = y + i * each
        h = each - 130000
        rrect(sl, MX, by, SW - 2 * MX, h, fill=C_LGRAY, radius=0.08)
        rrect(sl, MX, by, 70000, h, fill=C_GREEN, radius=0.5)
        cd = 300000
        oval(sl, MX + 180000, by + (h - cd) // 2, cd, cd, fill=C_BLUE)
        tb(sl, MX + 180000, by + (h - cd) // 2, cd, cd, [[("✓", {"font": F_MED, "sz": 14, "color": C_WHITE, "bold": True})]],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(sl, MX + 180000 + cd + 180000, by, SW - 2 * MX - 180000 - cd - 360000, h,
           [[(p, {"font": F_REG, "sz": 13})]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    footer(sl, titulo_aula, idx)
    return sl


def build_closing(prs, msg, sub):
    sl = add_slide(prs)
    topbar(sl)
    iw, ih = Image.open(SWIRL).size
    sw_w = 3800000
    swirl(sl, sw_w, (SW - sw_w) // 2, (SH - int(sw_w / (iw / ih))) // 2 - 300000)
    tb(sl, 0, SH // 2 - 700000, SW, 800000, [[(msg, {"font": F_MED, "sz": 44})]], align=PP_ALIGN.CENTER)
    tb(sl, 0, SH // 2 + 100000, SW, 400000, [[(sub, {"font": F_LIGHT, "sz": 16, "color": C_GRAY, "italic": True})]],
       align=PP_ALIGN.CENTER)
    iw, ih = Image.open(ADA_LOGO).size
    lw = 1000000; lh = int(lw / (iw / ih))
    pic(sl, ADA_LOGO, (SW - lw) // 2, SH - lh - 500000, lw, lh)
    tb(sl, 0, SH - 340000, SW, 240000, [[("Prof.: " + PROFESSOR, {"font": F_MED, "sz": 10, "color": C_GRAY})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return sl


# =============================================================================
#  CONTEÚDO DAS AULAS
# =============================================================================
IMG_SUPPLY = "supply-chain.png"
IMG_OPA = "desacoplamento-opa.png"
FONTE_SUPPLY = "clavis.com.br — Ameaças à cadeia de software"
FONTE_OPA = "Material do aluno (figura criada por IA)"

AULAS = [
    # ─── AULA 1 ───────────────────────────────────────────────────────────
    {"num": 1, "data": "Sexta, 16/10", "titulo": "Policy as Code e Compliance Automatizado",
     "subtitulo": "Como transformar regras de segurança em código versionável",
     "slides": [
         {"type": "image", "kicker": "Onde estamos", "titulo": "A cadeia de suprimentos de software", "image": IMG_SUPPLY,
          "fonte": FONTE_SUPPLY,
          "bullets": [("Desenvolvimento:", "código autoral, revisão, segredos"),
                      ("Dependências e build:", "bibliotecas, imagens, pipeline"),
                      ("Distribuição e deploy:", "registries, clusters, runtime"),
                      ("Este módulo:", "política → inventário → integridade → compliance")],
          "notas": "Situe o módulo como extensão do DevSecOps (CY-SE-010). Percorra a figura da esquerda para a direita e diga que cada elo terá um controle ao longo das 8 aulas. Pergunte onde a turma acha que o risco é maior.", "min": 8},
         {"type": "content", "kicker": "Problema gerador", "titulo": "Regra no papel, risco em produção",
          "bullets": [("A regra existe:", "\"nenhuma imagem roda como root\""),
                      ("A verificação:", "checklist manual, uma vez por trimestre"),
                      ("O resultado:", "17 imagens rodando como root em produção"),
                      ("O diagnóstico:", "o problema não é a política, é como ela é verificada")],
          "analogia": "Uma placa de 'use o cinto' sem sensor no banco: todo mundo conhece a regra, ninguém verifica. Policy as Code é o sensor que apita antes de o carro sair da garagem.",
          "notas": "Conte o caso das 17 imagens como falha de feedback, não de intenção. Pergunte: em que momento a equipe deveria ter descoberto? A resposta leva ao conceito de verificação automática em cada mudança.", "min": 10},
         {"type": "flow", "kicker": "Policy as Code", "titulo": "Da regra escrita ao código executável",
          "steps": [("Política", "Requisito de segurança ou compliance em linguagem natural"),
                    ("Código Rego", "Regra formal, versionada no Git e revisada por PR"),
                    ("Gate no pipeline", "Avaliada automaticamente a cada commit, build e deploy"),
                    ("Evidência", "Log do CI e decisão registrada — prova para auditoria")],
          "notas": "Explique que a mesma filosofia transformou infraestrutura (IaC) e pipelines (Pipeline as Code). Destaque os três pilares: código, automação e governança.", "min": 10},
         {"type": "comparison", "kicker": "Mudança de paradigma", "titulo": "Compliance manual vs. Policy as Code",
          "a_title": "Compliance manual", "a_items": ["Checklist em PDF ou Confluence", "Auditoria trimestral por humano", "Evidência em planilha", "Feedback semanas depois do deploy"],
          "b_title": "Policy as Code", "b_items": ["Regra em arquivo .rego versionado", "Avaliação automática por commit", "Evidência em código + logs do CI", "Feedback em segundos, no PR"],
          "notas": "Compare frequência, custo e qualidade da evidência. Reforce que a automação não elimina governança: exceções e aprovações continuam existindo, mas ficam rastreáveis.", "min": 10},
         {"type": "image", "kicker": "Princípio central", "titulo": "Decoupling: a aplicação pergunta, o OPA decide", "image": IMG_OPA,
          "fonte": FONTE_OPA,
          "bullets": [("Input:", "JSON que descreve o evento avaliado"),
                      ("Data:", "contexto externo, ex.: registries aprovados"),
                      ("Política:", "regras Rego mantidas pelo time de segurança"),
                      ("Decisão:", "allow / deny com mensagens explicativas")],
          "notas": "Mostre o fluxo no diagrama: a aplicação envia input, o OPA avalia contra políticas e dados e devolve a decisão. Vantagens: atualizar política sem redeploy, fonte única de verdade, reutilização em CI, Kubernetes e APIs.", "min": 12},
         {"type": "content", "kicker": "Supply chain na prática", "titulo": "Incidentes que motivam o paradigma",
          "bullets": [("SolarWinds (2020):", "build comprometido → exigir atestação de procedência"),
                      ("Codecov (2021):", "script adulterado → verificar hash de scripts de terceiros"),
                      ("Log4Shell (2021):", "CVE ubíqua → SBOM atualizado + gate por severidade"),
                      ("ua-parser-js (2021):", "conta sequestrada → alertar mudança de publicador"),
                      ("XZ Utils (2024):", "mantenedor malicioso → revisão multi-pessoa e build reprodutível")],
          "notas": "Para cada caso, pergunte: que política teria bloqueado ou detectado? Não entre em detalhes técnicos ainda — a Aula 6 aprofunda a anatomia dos ataques.", "min": 12},
         {"type": "flow", "kicker": "Security gates", "titulo": "Onde a política atua no pipeline",
          "steps": [("Commit / PR", "Secrets, licenças — pre-commit e Conftest"),
                    ("Build", "Dockerfile, IaC, dependências — Conftest no CI"),
                    ("Artefato", "SBOM, assinatura, atestações — Conftest + Cosign"),
                    ("Deploy", "Configuração de runtime — OPA Gatekeeper")],
          "analogia": "Pense em aeroporto: check-in, raio-x, portão e porta do avião. Cada gate verifica algo diferente e ninguém embarca pulando etapas.",
          "notas": "Conecte com os gates do módulo DevSecOps. Cada gate consome uma família de políticas e o Gatekeeper é a última barreira em runtime (tema da Aula 3).", "min": 10},
         {"type": "code", "kicker": "Primeiro contato", "titulo": "Uma política Rego em 8 linhas", "lang": "rego",
          "intro": "Em vez de um documento dizendo 'imagens não podem rodar como root':",
          "code": ["package docker.security", "",
                   "deny[msg] {", "    input.config.User == \"root\"",
                   "    msg := \"Imagem não pode executar como root.\"", "}", "",
                   "deny[msg] {", "    not input.config.User",
                   "    msg := \"Defina um usuário não-privilegiado.\"", "}"],
          "notas": "Leia a regra em voz alta: 'negue com esta mensagem quando o usuário for root'. Diga que a Aula 2 destrincha a sintaxe; hoje o objetivo é ver que uma política cabe em poucas linhas e é testável.", "min": 10},
         {"type": "code", "kicker": "Ferramenta do dia a dia", "titulo": "Conftest: OPA aplicado a arquivos reais",
          "intro": "O Conftest valida Dockerfiles, manifests Kubernetes, Terraform, JSON e YAML no pipeline:",
          "code": ["conftest test Dockerfile --policy ./policies/docker/",
                   "conftest test deployment.yaml --policy ./policies/kubernetes/",
                   "conftest test sbom.json --policy ./policies/sbom/", "",
                   "# Saída típica", "> FAIL - Dockerfile - Imagem não pode executar como root",
                   "> FAIL - Dockerfile - Base 'ubuntu:latest' não está na allowlist",
                   "> WARN - Dockerfile - Considere usar imagem distroless", "",
                   "> 2 tests, 0 passed, 1 warning, 2 failures"],
          "notas": "Mostre que o mesmo motor (OPA) atende CI (Conftest), Kubernetes (Gatekeeper) e aplicações (API). Cite Kyverno e Sentinel como alternativas do ecossistema.", "min": 8},
         {"type": "lab", "titulo": "Laboratório: primeira política e validação do ambiente",
          "passos": [("Instalar", "OPA e Conftest; validar com opa version e conftest --version"),
                     ("Escrever", "policies/docker/security.rego com dois deny"),
                     ("Avaliar", "opa eval com input root → violação; UID 1001 → vazio"),
                     ("Validar", "Dockerfile real com conftest test Dockerfile --policy policies/"),
                     ("Refletir", "como isso vira gate obrigatório em cada push?")],
          "entregavel": "Política Rego, dois inputs JSON, Dockerfile de teste e as saídas dos comandos. Runbook: labs/Aula_01_Runbook.md",
          "notas": "Confirme instalações antes de começar. Peça que prevejam a saída antes de cada comando. Circule entre os grupos e cuide para que as mensagens de deny sejam acionáveis.", "min": 60},
         {"type": "recap", "pontos": ["Policy as Code: políticas versionadas, testadas e executadas automaticamente a cada mudança.",
                                      "Decoupling: a aplicação pergunta, o OPA decide — uma política vale para CI, Kubernetes e APIs.",
                                      "Compliance contínuo substitui o checklist trimestral por evidência viva no pipeline.",
                                      "Conftest leva o OPA para Dockerfiles, manifests e SBOMs no CI."],
          "notas": "Retome a enquete inicial: política no papel vs. política executável. Antecipe a Aula 2: dominar a sintaxe Rego para escrever regras reais.", "min": 8},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: Linguagem Rego e Escrita de Políticas",
          "notas": "Encerre lembrando a atividade do LMS e a instalação de OPA/Conftest para quem ainda não concluiu.", "min": 3},
     ]},

    # ─── AULA 2 ───────────────────────────────────────────────────────────
    {"num": 2, "data": "Segunda, 19/10", "titulo": "Linguagem Rego e Escrita de Políticas",
     "subtitulo": "Como escrever uma política que reprova uma configuração insegura",
     "slides": [
         {"type": "content", "kicker": "Modelo mental", "titulo": "Declarativo: descreva condições, não passos",
          "bullets": [("Inspirada em Datalog:", "você descreve quando algo é verdadeiro"),
                      ("Opera sobre JSON:", "input, data e políticas são estruturas de dados"),
                      ("Condições com AND implícito:", "todas verdadeiras → regra produz valor"),
                      ("Resultado:", "políticas concisas, auditáveis e testáveis")],
          "analogia": "É como um filtro de e-mail: você não programa 'como procurar spam', você descreve como o spam se parece — 'se vier de X e contém Y, marque'.",
          "notas": "Contraste com código imperativo. O desconforto inicial é normal; a fluência vem repetindo os exemplos. Avise que hoje a aula é majoritariamente prática.", "min": 10},
         {"type": "flow", "kicker": "Anatomia", "titulo": "Os quatro blocos de uma política",
          "steps": [("package", "Namespace: data.docker.security"),
                    ("Regras", "deny[msg] { ... } ou allow = true { ... }"),
                    ("input / data", "Evento avaliado e contexto externo (allowlists)"),
                    ("Decisão", "Conjunto de mensagens ou valor único")],
          "notas": "Percorra o fluxo input → regra → query → decisão. Diferencie input (muda a cada avaliação) de data (carregado como configuração).", "min": 10},
         {"type": "comparison", "kicker": "Tipos de regra", "titulo": "Regra completa vs. regra parcial",
          "a_title": "Completa — um valor", "a_items": ["default allow = false", "allow { input.user.role == \"admin\" }", "Resultado escalar: true / false", "Uso: decisões binárias, flags"],
          "b_title": "Parcial — um conjunto", "b_items": ["deny[msg] { ... }", "Cada bloco que casar adiciona uma msg", "Resultado: lista de violações", "Uso: deny, warn, violation"],
          "notas": "Mostre que vários blocos deny com o mesmo nome se somam. 'default allow = false' é falha segura: se nada casar, nega.", "min": 10},
         {"type": "code", "kicker": "Sintaxe essencial", "titulo": "Variáveis, iteração e built-ins", "lang": "rego",
          "intro": ":= declara, == compara, _ itera; sprintf monta a mensagem:",
          "code": ["deny[msg] {", "    container := input.spec.containers[_]",
                   "    container.image == \"nginx:latest\"",
                   "    msg := sprintf(\"Container '%v' usa tag latest.\", [container.name])", "}", "",
                   "deny[msg] {", "    image := input.spec.containers[_].image",
                   "    not regex.match(`.*@sha256:[a-f0-9]{64}$`, image)",
                   "    msg := sprintf(\"Imagem '%v' deve usar digest.\", [image])", "}"],
          "notas": "Explique _ como 'para cada'. Cite as famílias de built-ins: strings (startswith, contains), coleções (count), regex, conjuntos (in). Não decorar — consultar a referência.", "min": 12},
         {"type": "code", "kicker": "Dockerfile", "titulo": "Política para Dockerfile (formato Conftest)", "lang": "rego",
          "intro": "O Conftest converte cada linha em {Cmd, Value}:",
          "code": ["package main", "",
                   "approved_base := {\"cgr.dev/chainguard/static\", \"gcr.io/distroless/base\"}", "",
                   "deny[msg] {", "    input[i].Cmd == \"from\"", "    image := input[i].Value[0]",
                   "    not is_approved(image)",
                   "    msg := sprintf(\"Linha %d: base '%v' não aprovada.\", [i, image])", "}", "",
                   "deny[msg] { not has_user; msg := \"Dockerfile sem USER.\" }",
                   "has_user { input[_].Cmd == \"user\" }",
                   "is_approved(img) { startswith(img, approved_base[_]) }"],
          "notas": "Destaque a função helper is_approved e a allowlist como conjunto. Pergunte o que acontece com 'FROM ubuntu:latest'.", "min": 12},
         {"type": "code", "kicker": "Kubernetes", "titulo": "Política para manifests Kubernetes", "lang": "rego",
          "intro": "YAML vira JSON; o padrão é o mesmo:",
          "code": ["package kubernetes.pod_security", "",
                   "deny[msg] {", "    input.kind == \"Deployment\"",
                   "    c := input.spec.template.spec.containers[_]",
                   "    c.securityContext.runAsUser == 0",
                   "    msg := sprintf(\"'%v' não pode rodar como UID 0.\", [c.name])", "}", "",
                   "deny[msg] {", "    c := input.spec.template.spec.containers[_]",
                   "    not c.resources.limits.memory",
                   "    msg := sprintf(\"'%v' sem limite de memória.\", [c.name])", "}"],
          "notas": "Mostre o caminho spec.template.spec.containers no YAML real. Mencione a política de secrets em env vars como exercício extra.", "min": 10},
         {"type": "code", "kicker": "Qualidade", "titulo": "Testes com opa test", "lang": "rego",
          "intro": "Arquivo *_test.rego, regras test_*, 'with input as':",
          "code": ["package main", "",
                   "test_deny_root_user {",
                   "    deny[\"Linha 2: USER 'root' proibido.\"] with input as [",
                   "        {\"Cmd\": \"from\", \"Value\": [\"cgr.dev/chainguard/static:latest\"]},",
                   "        {\"Cmd\": \"user\", \"Value\": [\"root\"]}", "    ]", "}", "",
                   "test_allow_valid { count(deny) == 0 with input as [...] }", "",
                   "# opa test policies/ -v", "> PASS: 3/3",
                   "# opa test policies/ --coverage --format=json"],
          "notas": "Toda política de produção tem teste positivo e negativo. Mostre a cobertura como forma de achar caminhos não testados.", "min": 10},
         {"type": "content", "kicker": "Padrões", "titulo": "Boas práticas que evitam dor de cabeça",
          "bullets": [("Um pacote por domínio:", "docker/, kubernetes/, sbom/, shared/"),
                      ("Mensagens acionáveis:", "diga exatamente o que corrigir"),
                      ("Allowlist em data:", "atualize listas sem tocar na política"),
                      ("Exceções documentadas:", "ticket e prazo no próprio código"),
                      ("Revisão por PR:", "mesmo rigor do código de produção")],
          "analogia": "Política sem teste e sem comentário é como regra de condomínio colada no elevador sem assinatura: ninguém sabe quem decidiu, por quê, nem se ainda vale.",
          "notas": "Mostre a estrutura de repositório recomendada. Reforce: política é código de produção, com dono, teste e histórico.", "min": 8},
         {"type": "lab", "titulo": "Laboratório: política completa para Dockerfile e Kubernetes",
          "passos": [("Estruturar", "policies/docker, policies/kubernetes, manifests/"),
                     ("Docker:", "allowlist de base, bloquear USER root, exigir USER"),
                     ("Quebrar:", "trocar para ubuntu:latest + root e ver as violações"),
                     ("Kubernetes:", "bloquear UID 0 e exigir limite de memória"),
                     ("Testar:", "security_test.rego → opa test policies/ -v")],
          "entregavel": "Políticas Docker/Kubernetes, manifests de teste e *_test.rego passando. Runbook: labs/Aula_02_Runbook.md",
          "notas": "Comece pelo Dockerfile e só depois Kubernetes. Peça que cada grupo quebre a própria política de propósito para ver a mensagem de erro.", "min": 70},
         {"type": "recap", "pontos": ["Rego é declarativo: descreva as condições que tornam a violação verdadeira.",
                                      "deny[msg] coleta violações; default allow = false garante falha segura.",
                                      "Conftest aplica a mesma política a Dockerfiles, YAML e JSON reais.",
                                      "opa test com casos positivos e negativos previne regressões."],
          "notas": "Peça a alguém para explicar quando usar regra parcial vs. completa. Ponte: a Aula 3 impõe essas políticas no cluster.", "min": 5},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: Enforcement com OPA Gatekeeper",
          "notas": "Peça que tragam kind/minikube e kubectl instalados para a próxima aula.", "min": 3},
     ]},

    # ─── AULA 3 ───────────────────────────────────────────────────────────
    {"num": 3, "data": "Terça, 20/10", "titulo": "Enforcement: OPA Gatekeeper e Admission Control",
     "subtitulo": "Como impor políticas de forma centralizada em um cluster",
     "slides": [
         {"type": "content", "kicker": "Problema gerador", "titulo": "Política decorativa vs. enforcement",
          "bullets": [("Relatório:", "política que só gera visibilidade"),
                      ("Enforcement:", "política que muda o resultado da operação"),
                      ("Gate obrigatório:", "build inseguro não progride, manifest é rejeitado"),
                      ("Exceções:", "documentadas, aprovadas e auditáveis")],
          "analogia": "Um radar que só fotografa não reduz velocidade. O que muda o comportamento é a multa chegando em casa — e, melhor ainda, a lombada antes da escola.",
          "notas": "Pergunte: um desenvolvedor consegue contornar o CI e aplicar direto no cluster? Essa possibilidade justifica a segunda barreira em runtime.", "min": 8},
         {"type": "comparison", "kicker": "Dois mundos", "titulo": "Enforcement em CI/CD vs. em runtime",
          "a_title": "CI/CD — primeira linha", "a_items": ["Antes do merge ou deploy", "Conftest / OPA CLI no pipeline", "Bloqueia PR, build, deploy", "Dev corrige no mesmo PR"],
          "b_title": "Runtime — última barreira", "b_items": ["Quando o recurso é criado no cluster", "OPA Gatekeeper / Kyverno", "Bloqueia criação ou atualização", "Protege contra bypass do pipeline"],
          "notas": "Boa prática é combinar os dois: feedback rápido no CI e garantia no cluster. Não são alternativas, são camadas.", "min": 10},
         {"type": "flow", "kicker": "Admission Control", "titulo": "O caminho de um kubectl apply",
          "steps": [("kubectl apply", "Usuário envia o manifest"),
                    ("API Server", "Autentica, autoriza e valida sintaxe"),
                    ("Webhook Gatekeeper", "Avalia todas as Constraints aplicáveis"),
                    ("Admitido / Negado", "deny rejeita; permitido persiste no etcd")],
          "notas": "Ressalte que o recurso só é gravado depois da admissão. A resposta do webhook volta como erro do kubectl, com a mensagem da política.", "min": 10},
         {"type": "image", "kicker": "O motor por trás", "titulo": "Gatekeeper = OPA dentro do cluster", "image": IMG_OPA,
          "fonte": FONTE_OPA,
          "bullets": [("Mesmo OPA:", "input agora é o AdmissionReview do Kubernetes"),
                      ("CRDs:", "políticas viram recursos versionáveis do cluster"),
                      ("ConstraintTemplate:", "define a política Rego e seus parâmetros"),
                      ("Constraint:", "instancia com escopo, parâmetros e ação")],
          "notas": "Relacione com o diagrama de desacoplamento da Aula 1: a 'aplicação' agora é o API Server. Template é reutilizável; Constraint é a instância.", "min": 10},
         {"type": "code", "kicker": "ConstraintTemplate", "titulo": "Rego dentro de um CRD", "lang": "yaml",
          "intro": "Define o tipo K8sPSPRunAsNonRoot e a regra violation:",
          "code": ["apiVersion: templates.gatekeeper.sh/v1", "kind: ConstraintTemplate",
                   "metadata: { name: k8spsprunasnonroot }", "spec:",
                   "  crd: { spec: { names: { kind: K8sPSPRunAsNonRoot } } }",
                   "  targets:", "    - target: admission.k8s.gatekeeper.sh", "      rego: |",
                   "        package k8spsprunasnonroot",
                   "        violation[{\"msg\": msg}] {",
                   "            c := input.review.object.spec.containers[_]",
                   "            not c.securityContext.runAsNonRoot",
                   "            msg := sprintf(\"'%v' deve declarar runAsNonRoot.\", [c.name])",
                   "        }"],
          "notas": "Aponte input.review.object: é o recurso sendo criado. A regra se chama violation, não deny — convenção do Gatekeeper.", "min": 10},
         {"type": "code", "kicker": "Constraint", "titulo": "Escopo, parâmetros e ação", "lang": "yaml",
          "intro": "Aplica o template a Pods de produção e staging, bloqueando:",
          "code": ["apiVersion: constraints.gatekeeper.sh/v1beta1", "kind: K8sPSPRunAsNonRoot",
                   "metadata: { name: pods-must-run-as-non-root }", "spec:",
                   "  enforcementAction: deny      # dryrun | warn | deny", "  match:",
                   "    kinds: [{ apiGroups: [\"\"], kinds: [\"Pod\"] }]",
                   "    namespaces: [production, staging]",
                   "    excludedNamespaces: [kube-system, gatekeeper-system]",
                   "  parameters:", "    exemptImages: [\"gcr.io/google-containers/pause\"]"],
          "notas": "Mostre os três pontos de controle: enforcementAction, match e parameters. Exceção via parâmetro é a mais auditável porque fica explícita no recurso.", "min": 10},
         {"type": "flow", "kicker": "Rollout seguro", "titulo": "dryrun → warn → deny → audit",
          "steps": [("dryrun", "Avalia e registra em status, sem bloquear. 2 a 4 semanas."),
                    ("warn", "Permite, mas avisa o cliente. Onboarding gradual."),
                    ("deny", "Rejeita o recurso. Só após cobrir exceções."),
                    ("audit", "Varre recursos existentes periodicamente: status.violations")],
          "analogia": "Como mudar o limite de velocidade de uma rua: primeiro placa educativa, depois aviso, depois multa — e o radar fixo continua medindo quem já estava lá.",
          "notas": "Explique por que começar em dryrun: medir impacto, comunicar times, dar prazo. Audit pega recursos criados antes da política ou via bypass.", "min": 10},
         {"type": "code", "kicker": "Pipeline", "titulo": "Gate obrigatório no GitHub Actions", "lang": "yaml",
          "intro": "Conftest + opa test rodando em todo pull request:",
          "code": ["name: Policy Gate", "on: { pull_request: { branches: [main] } }", "jobs:",
                   "  policy-check:", "    runs-on: ubuntu-latest", "    steps:",
                   "      - uses: actions/checkout@b4ffde65f46336ab88eb53be808477a3936bae11",
                   "      - name: Conftest em Dockerfiles",
                   "        run: conftest test Dockerfile --policy policies/docker/",
                   "      - name: Conftest em manifests",
                   "        run: conftest test k8s/*.yaml --policy policies/kubernetes/",
                   "      - name: Testes das políticas", "        run: opa test policies/ -v"],
          "notas": "Chame atenção para a action fixada por SHA de commit, não por tag — isso é supply chain security do próprio pipeline (tema da Aula 6). Métricas Prometheus do Gatekeeper alimentam dashboards.", "min": 8},
         {"type": "lab", "titulo": "Laboratório: do Conftest ao Gatekeeper",
          "passos": [("Cluster:", "kind create cluster --name policy-lab"),
                     ("Instalar:", "Gatekeeper e aguardar pods em gatekeeper-system"),
                     ("Aplicar:", "ConstraintTemplate + Constraint (dryrun → deny)"),
                     ("Testar:", "Pod sem securityContext → Forbidden; corrigir → criado"),
                     ("Inspecionar:", "kubectl get constraint ... -o yaml → status.violations")],
          "entregavel": "Template, Constraint, manifests inseguro/seguro e a mensagem de bloqueio capturada. Runbook: labs/Aula_03_Runbook.md",
          "notas": "Verifique kind/kubectl antes. Peça que prevejam a decisão para o Pod sem securityContext e observem a mensagem do webhook.", "min": 70},
         {"type": "recap", "pontos": ["Enforcement transforma política em gate: CI/CD dá feedback cedo, Gatekeeper garante no cluster.",
                                      "ConstraintTemplate define a regra; Constraint define escopo, parâmetros e ação.",
                                      "Rollout gradual: dryrun → warn → deny, com exceções documentadas.",
                                      "Audit contínuo encontra recursos não-conformes já existentes."],
          "notas": "Pergunta de fechamento: CI substitui admission? Resposta esperada: não, são camadas. Ponte para SBOM.", "min": 6},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: SBOM — Fundamentos, SPDX e CycloneDX",
          "notas": "Peça Syft, Trivy e jq instalados para a próxima aula.", "min": 3},
     ]},

    # ─── AULA 4 ───────────────────────────────────────────────────────────
    {"num": 4, "data": "Sexta, 23/10", "titulo": "SBOM: Fundamentos, SPDX e CycloneDX",
     "subtitulo": "Você sabe exatamente quais componentes compõem seu software?",
     "slides": [
         {"type": "content", "kicker": "Problema gerador", "titulo": "Dezembro de 2021: \"nós usamos Log4j?\"",
          "bullets": [("A pergunta:", "em quais aplicações? em quais versões?"),
                      ("Sem inventário:", "dias de grep em repositórios"),
                      ("Com SBOM:", "resposta em segundos"),
                      ("Custo da lentidão:", "incidentes, exposição e remediação emergencial")],
          "analogia": "SBOM é a lista de ingredientes do rótulo. Quando sai um recall de um ingrediente, quem tem rótulo sabe na hora quais produtos tirar da prateleira.",
          "notas": "Use a pergunta 'como sua equipe responderia hoje a uma nova CVE?' como gancho. Tempo de resposta é proporcional à qualidade do inventário.", "min": 8},
         {"type": "image", "kicker": "Onde o SBOM vive", "titulo": "O inventário acompanha cada elo da cadeia", "image": IMG_SUPPLY,
          "fonte": FONTE_SUPPLY,
          "bullets": [("Dependências:", "typosquatting, CVEs → SCA + SBOM"),
                      ("Build:", "SBOM gerado junto com o artefato"),
                      ("Armazenamento:", "SBOM anexado à imagem no registry"),
                      ("Operação:", "novas CVEs cruzadas com SBOM continuamente")],
          "notas": "Mostre que o SBOM nasce no build e é consumido até a operação. OWASP A03 — Software Supply Chain Failures — é o tema central desta e das próximas duas aulas.", "min": 8},
         {"type": "content", "kicker": "Definição", "titulo": "Os 7 elementos mínimos (NTIA)",
          "bullets": [("Fornecedor:", "quem criou ou distribui o componente"),
                      ("Nome e versão:", "identificação exata"),
                      ("Identificador único:", "PURL, CPE ou SWID"),
                      ("Relações de dependência:", "quem depende de quem"),
                      ("Autor e timestamp:", "quem gerou o SBOM e quando")],
          "analogia": "É uma nota fiscal do software: cada item, sua origem, seu código de barras e a data da emissão.",
          "notas": "Cite os motores regulatórios: Executive Order 14028 (EUA) e o EU Cyber Resilience Act. Além de compliance, há valor operacional direto.", "min": 10},
         {"type": "comparison", "kicker": "Dois formatos abertos", "titulo": "SPDX vs. CycloneDX",
          "a_title": "SPDX — Linux Foundation", "a_items": ["Criado em 2010 · ISO/IEC 5962:2021", "Foco original: licenças e compliance legal", "JSON, YAML, RDF, tag-value", "Forte em auditoria jurídica"],
          "b_title": "CycloneDX — OWASP", "b_items": ["Criado em 2017 · ECMA TC54", "Foco original: segurança e SCA", "JSON, XML, Protobuf", "VEX nativo, ML-BOM, atestações"],
          "notas": "Não é escolha binária: muitas organizações geram nos dois. SPDX para o jurídico, CycloneDX para a operação de segurança.", "min": 10},
         {"type": "code", "kicker": "Identificador universal", "titulo": "PURL: o ponto de correlação com CVEs",
          "intro": "pkg:<type>/<namespace>/<name>@<version> — o mesmo em SPDX e CycloneDX:",
          "code": ["pkg:npm/express@4.19.2", "pkg:pypi/django@5.0.6",
                   "pkg:maven/org.apache.logging.log4j/log4j-core@2.17.1",
                   "pkg:docker/nginx@1.27.0", "pkg:golang/github.com/gin-gonic/gin@v1.10.0", "",
                   "# Trecho CycloneDX", "{ \"type\": \"library\", \"name\": \"express\", \"version\": \"4.19.2\",",
                   "  \"purl\": \"pkg:npm/express@4.19.2\",",
                   "  \"licenses\": [{\"license\": {\"id\": \"MIT\"}}],",
                   "  \"hashes\": [{\"alg\": \"SHA-256\", \"content\": \"a1b2c3...\"}] }"],
          "notas": "PURL é a chave que plataformas usam para consultar OSV, NVD e GitHub Advisory. Sem PURL, não há correlação automática.", "min": 8},
         {"type": "code", "kicker": "Ferramentas", "titulo": "Gerando SBOM com Syft e Trivy",
          "intro": "Syft detecta npm, pip, maven, go, apk, deb, rpm… automaticamente:",
          "code": ["# Syft — imagem Docker, nos dois formatos",
                   "syft nginx:1.27.0 -o cyclonedx-json=nginx.cdx.json -o spdx-json=nginx.spdx.json", "",
                   "# Syft — diretório local", "syft dir:./minha-app -o spdx-json > app.spdx.json", "",
                   "# Trivy — SBOM e vulnerabilidades em um passo",
                   "trivy image --format cyclonedx --output nginx.cdx.json nginx:1.27.0", "",
                   "# Inspecionar", "jq '.components | length' nginx.cdx.json",
                   "jq -r '.components[].licenses[]?.license.id' nginx.cdx.json | sort -u"],
          "notas": "Cite cdxgen (OWASP) e microsoft/sbom-tool como alternativas. Mostre o jq como forma rápida de inspecionar antes de ir para uma plataforma.", "min": 10},
         {"type": "code", "kicker": "Pipeline", "titulo": "SBOM como artefato obrigatório do release", "lang": "yaml",
          "intro": "Gerado no build e versionado junto com a imagem:",
          "code": ["- name: Build", "  run: docker build -t minha-api:${{ github.sha }} .", "",
                   "- name: SBOM (CycloneDX)", "  uses: anchore/sbom-action@v0",
                   "  with: { image: minha-api:${{ github.sha }}, format: cyclonedx-json, output-file: sbom.cdx.json }", "",
                   "- name: SBOM (SPDX)", "  uses: anchore/sbom-action@v0",
                   "  with: { image: minha-api:${{ github.sha }}, format: spdx-json, output-file: sbom.spdx.json }", "",
                   "- uses: actions/upload-artifact@v4", "  with: { name: sboms, path: sbom.*.json }"],
          "notas": "O SBOM precisa acompanhar o release: mesmo commit, mesma imagem. Na Aula 7 ele será anexado à imagem como atestação assinada.", "min": 8},
         {"type": "code", "kicker": "Conectando com Policy as Code", "titulo": "Política Rego sobre o próprio SBOM", "lang": "rego",
          "intro": "Validar o inventário antes de aceitar o artefato:",
          "code": ["package main", "forbidden := {\"GPL-3.0-only\", \"AGPL-3.0-only\"}", "",
                   "deny[msg] {", "    c := input.components[_]", "    lic := c.licenses[_].license.id",
                   "    forbidden[lic]",
                   "    msg := sprintf(\"'%v' com licença não permitida: %v.\", [c.name, lic])", "}", "",
                   "deny[msg] {", "    c := input.components[_]", "    not c.purl",
                   "    msg := sprintf(\"'%v' sem PURL.\", [c.name])", "}", "",
                   "# conftest test sbom.cdx.json --policy policies/sbom/"],
          "notas": "Fecha o ciclo com as Aulas 1–3: a mesma linguagem valida Dockerfile, Kubernetes e SBOM. Mencione AIBOM/ML-BOM: modelos e datasets também entram no inventário (CycloneDX 1.5).", "min": 10},
         {"type": "lab", "titulo": "Laboratório: gerar e analisar SBOMs",
          "passos": [("Gerar:", "nginx.cdx.json e nginx.spdx.json com Syft; repetir com Trivy"),
                     ("Inspecionar:", "contar componentes, listar versões e licenças com jq"),
                     ("Comparar:", "campos de um componente em SPDX vs. CycloneDX"),
                     ("Validar:", "política Rego de licença + PURL via Conftest"),
                     ("Automatizar:", "Makefile com sbom-cdx, sbom-spdx e policy")],
          "entregavel": "Dois SBOMs, consultas jq e política de validação com resultado. Runbook: labs/Aula_04_Runbook.md",
          "notas": "Diferenças de contagem entre formatos e ferramentas são normais — discuta, não trate como erro. Guarde os SBOMs: a Aula 5 os ingere no Dependency-Track.", "min": 75},
         {"type": "recap", "pontos": ["SBOM é o inventário legível por máquina: fornecedor, versão, PURL, relações, autor, timestamp.",
                                      "SPDX (ISO, licenças) e CycloneDX (OWASP, segurança, VEX, ML-BOM) são complementares.",
                                      "Syft e Trivy geram SBOM no build; o artefato acompanha o release.",
                                      "Policy as Code valida o SBOM: licenças proibidas e componentes sem PURL."],
          "notas": "Pergunta de fechamento: SBOM é inventário ou scanner? Inventário. A correlação com risco é a próxima aula.", "min": 5},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: Análise de SBOM e Gestão de Dependências",
          "notas": "Peça Docker Compose funcionando para subir o Dependency-Track na próxima aula.", "min": 3},
     ]},

    # ─── AULA 5 ───────────────────────────────────────────────────────────
    {"num": 5, "data": "Segunda, 26/10", "titulo": "Análise de SBOM e Gestão de Dependências",
     "subtitulo": "Como descobrir em minutos se você é afetado por uma nova CVE",
     "slides": [
         {"type": "content", "kicker": "Problema gerador", "titulo": "Log4Shell em segundos ou em dias?",
          "bullets": [("Gestão imatura:", "dias de grep em repositórios"),
                      ("Gestão madura:", "dashboard lista apps afetadas e versões"),
                      ("A diferença:", "não são os engenheiros, é a infraestrutura de governança"),
                      ("Hoje:", "transformar inventário em decisão e remediação")],
          "analogia": "Ter a lista de ingredientes é o primeiro passo; o que salva é o sistema que cruza o recall com todos os produtos e já aciona a retirada.",
          "notas": "Retome o SBOM da Aula 4: gerar é metade. Hoje o valor surge quando uma plataforma correlaciona componentes com bases de vulnerabilidade e dispara correções.", "min": 8},
         {"type": "flow", "kicker": "Ciclo fechado", "titulo": "Gestão contínua de dependências",
          "steps": [("SBOM", "Gerado em cada build"),
                    ("Dependency-Track", "Correlaciona PURLs com NVD, OSV, GitHub Advisory"),
                    ("Policy / alerta", "FAIL para CVE crítica; Slack, Jira, webhook"),
                    ("PR automático", "Dependabot ou Renovate propõem a correção"),
                    ("Pipeline valida", "Testes + políticas Rego; merge; novo SBOM")],
          "notas": "Enfatize o ciclo: a saída (novo build) realimenta a entrada (novo SBOM). De esforço reativo para processo preventivo.", "min": 10},
         {"type": "content", "kicker": "Fontes de verdade", "titulo": "Bases de vulnerabilidade",
          "bullets": [("NVD (NIST):", "CVEs com pontuação CVSS"),
                      ("OSV (Google):", "vulnerabilidades por ecossistema open source"),
                      ("GitHub Advisory Database:", "CVEs e advisories de pacotes"),
                      ("EPSS (FIRST):", "probabilidade real de exploração"),
                      ("Combinação:", "NVD + OSV + GHSA cobrem quase todo o espaço relevante")],
          "notas": "CVSS mede severidade técnica; EPSS mede probabilidade de exploração. Priorização boa combina os dois com contexto da aplicação.", "min": 10},
         {"type": "code", "kicker": "Dependency-Track", "titulo": "Ingestão via API",
          "intro": "Cada projeto é uma aplicação ou componente; autoCreate cria na primeira ingestão:",
          "code": ["export DT_URL=\"http://localhost:8081\"",
                   "export DT_KEY=\"<api-key em variável, nunca no Git>\"", "",
                   "curl -X POST \"${DT_URL}/api/v1/bom\" \\",
                   "  -H \"X-Api-Key: ${DT_KEY}\" \\",
                   "  -F \"autoCreate=true\" \\",
                   "  -F \"projectName=minha-api\" \\",
                   "  -F \"projectVersion=1.2.3\" \\",
                   "  -F \"bom=@sbom.cdx.json\"", "",
                   "# Em segundos: componentes, CVEs e licenças no dashboard"],
          "notas": "Mostre o docker-compose (apiserver + frontend). Reforce: API key em secret do CI. DefectDojo consolida findings de várias ferramentas (SAST, DAST, SCA) em uma visão única.", "min": 10},
         {"type": "content", "kicker": "Policies", "titulo": "Políticas de aceite e risco",
          "bullets": [("Sem CVE crítica:", "Severity == CRITICAL → FAIL"),
                      ("CVE alta envelhecida:", "HIGH e Age > 30 dias → FAIL"),
                      ("Licença proibida:", "GPL-3.0, AGPL-3.0 → FAIL"),
                      ("Componente desatualizado:", "última versão há > 2 anos → WARN"),
                      ("Ações:", "FAIL alerta e bloqueia; WARN avisa; INFO registra")],
          "analogia": "É o regulamento do porto: carga com contrabando não desembarca (FAIL), carga com documentação vencida recebe aviso (WARN), o resto só é registrado.",
          "notas": "Policy no Dependency-Track = condição + ação. Discuta exceções com prazo: aceitar risco é decisão documentada, não ausência de decisão.", "min": 10},
         {"type": "code", "kicker": "Gate", "titulo": "Bloqueando o pipeline pela API",
          "intro": "Consulta violações FAIL do projeto e falha o build:",
          "code": ["PROJECT_UUID=$(curl -s -H \"X-Api-Key: ${DT_KEY}\" \\",
                   "  \"${DT_URL}/api/v1/project/lookup?name=lab-node14&version=baseline\" \\",
                   "  | jq -r .uuid)", "",
                   "VIOLATIONS=$(curl -s -H \"X-Api-Key: ${DT_KEY}\" \\",
                   "  \"${DT_URL}/api/v1/violation/project/${PROJECT_UUID}\" \\",
                   "  | jq '[.[] | select(.policyCondition.policy.violationState == \"FAIL\")]'\\",
                   "  | jq length)", "",
                   "echo \"Violações: $VIOLATIONS\"",
                   "test \"$VIOLATIONS\" -eq 0 || { echo \"Pipeline bloqueado\"; exit 1; }"],
          "notas": "Leia o comando: lookup do projeto, lista de violações, filtro por FAIL, exit 1. É o mesmo padrão de gate das Aulas 1–3, agora alimentado pelo SBOM.", "min": 8},
         {"type": "comparison", "kicker": "Remediação automatizada", "titulo": "Dependabot vs. Renovate",
          "a_title": "Dependabot", "a_items": ["Nativo no GitHub", "Config YAML simples (.github/dependabot.yml)", "20+ ecossistemas", "Agrupamento e automerge configuráveis"],
          "b_title": "Renovate", "b_items": ["GitHub, GitLab, Bitbucket, Azure DevOps", "JSON com regras granulares", "90+ ecossistemas", "Automerge por tipo (patch/minor) e dashboard"],
          "notas": "Descobrir é metade; atualizar sem quebrar é a outra. Mostre o exemplo de automerge para patch/minor e aprovação manual para major.", "min": 8},
         {"type": "code", "kicker": "Python e IA", "titulo": "pip-audit e dependências de modelos",
          "intro": "Ecossistema científico tem particularidades — e modelos são dependências de primeira classe:",
          "code": ["pip install pip-audit", "pip-audit --strict", "",
                   "> Found 2 known vulnerabilities in 2 packages",
                   "> requests  2.25.0  GHSA-j8r2-6x86-q33q  Fix: 2.31.0",
                   "> urllib3   1.26.5  GHSA-v845-jxx5-vc9f  Fix: 1.26.18", "",
                   "# Dependências de IA que escapam do SCA tradicional:",
                   "#  pesos (.safetensors) → hash + assinatura + fonte verificada",
                   "#  datasets → SBOM estendido + hash",
                   "#  frameworks (PyTorch, TF) → SCA tradicional"],
          "notas": "OWASP Top 10 for LLM recomenda tratar modelos como dependências: procedência, hash e licença no SBOM. Lockfiles com hash (poetry, pipenv) são preferíveis a requirements.txt solto.", "min": 8},
         {"type": "lab", "titulo": "Laboratório: do SBOM ao gate de pipeline",
          "passos": [("Subir:", "Dependency-Track via Docker Compose; trocar senha; gerar API key"),
                     ("Gerar e enviar:", "syft node:14-alpine → POST /api/v1/bom (lab-node14)"),
                     ("Observar:", "vulnerabilidades correlacionadas no projeto"),
                     ("Policy:", "Severity is CRITICAL → FAIL; ver violações"),
                     ("Gate:", "script via API retorna exit 1; opcional: Dependabot e pip-audit")],
          "entregavel": "SBOM, evidência da ingestão, policy configurada e saída do gate. Runbook: labs/Aula_05_Runbook.md",
          "notas": "O backend demora alguns minutos para inicializar — use o tempo para a discussão de policies. API key só em variável de ambiente.", "min": 75},
         {"type": "recap", "pontos": ["Dependency-Track ingere SBOMs e correlaciona PURLs com NVD, OSV e GitHub Advisory.",
                                      "Policies (severidade, idade, licença) viram gate via API — descoberta sem remediação é só visibilidade.",
                                      "Dependabot e Renovate fecham o ciclo com PRs automáticos validados pelo pipeline.",
                                      "Python e IA: lockfiles com hash, pip-audit e modelos como dependências de primeira classe."],
          "notas": "Pergunta: CVSS basta para priorizar? Não — EPSS e contexto. Ponte: Aula 6 olha a cadeia pelos olhos do atacante.", "min": 5},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: Supply Chain Security — Ataques e Frameworks",
          "notas": "Sugira a leitura prévia dos casos SolarWinds e XZ Utils no material.", "min": 3},
     ]},

    # ─── AULA 6 ───────────────────────────────────────────────────────────
    {"num": 6, "data": "Terça, 27/10", "titulo": "Supply Chain Security: Ataques e Frameworks",
     "subtitulo": "Como um atacante compromete seu produto sem tocar no seu código",
     "slides": [
         {"type": "image", "kicker": "Modelo de ameaça", "titulo": "Três fases, três fronteiras de confiança", "image": IMG_SUPPLY,
          "fonte": FONTE_SUPPLY,
          "bullets": [("Origem (Source):", "commit malicioso, credenciais vazadas, branch sem proteção"),
                      ("Build:", "typosquatting, dependency confusion, build poisoning"),
                      ("Distribuição:", "artefato substituído, downgrade, chave roubada, mirror malicioso"),
                      ("Confiança transitiva:", "você confia em toda a infraestrutura que produziu o pacote")],
          "notas": "Use o modelo SLSA de três fases. Cada fase tem uma fronteira de confiança; sem isolamento, um comprometimento intermediário chega à produção sem rastro.", "min": 10},
         {"type": "content", "kicker": "Fase de build", "titulo": "Vetores que exploram dependências",
          "bullets": [("Typosquatting:", "reqests, urllib, colorz — erro de digitação no install"),
                      ("Dependency confusion:", "pacote público com nome do pacote interno e versão maior"),
                      ("Pacote malicioso:", "postinstall que exfiltra variáveis de ambiente"),
                      ("Build poisoning:", "toolchain ou plugin do CI injeta código na compilação"),
                      ("Cache envenenado:", "chave previsível + ACL fraca")],
          "analogia": "Dependency confusion é como um entregador que, vendo dois 'João Silva' no prédio, entrega para o que gritar mais alto — o instalador pega a versão 'mais nova', mesmo sendo do atacante.",
          "notas": "Para cada vetor, peça a mitigação: lockfile com hash, scope privado + registry autenticado, ignore-scripts, build efêmero com atestação, chave de cache com hash do lockfile.", "min": 12},
         {"type": "content", "kicker": "Estudos de caso", "titulo": "Quatro ataques, quatro lições",
          "bullets": [("SolarWinds / SUNBURST (2020):", "MSBuild injetava DLL; artefato saía assinado. 18 mil organizações."),
                      ("Codecov (2021):", "bash uploader exfiltrou secrets de CI por 73 dias"),
                      ("ua-parser-js (2021):", "conta npm sequestrada; 8 mi downloads/semana"),
                      ("XZ Utils (2024):", "2 anos de confiança construída; backdoor no sshd via tarball")],
          "analogia": "XZ Utils é o 'funcionário exemplar' que passa dois anos ganhando confiança para, um dia, deixar a porta dos fundos aberta — e só foi pego porque alguém notou 500 ms a mais no login.",
          "notas": "Extraia a lição defensiva de cada caso: build reprodutível e atestação (SolarWinds), hash de scripts de bootstrap (Codecov), 2FA e assinaturas (ua-parser-js), build do Git e não do tarball (XZ). Cite Kaseya e NotPetya como ataques em cascata.", "min": 15},
         {"type": "flow", "kicker": "Framework de maturidade", "titulo": "SLSA — Build track",
          "steps": [("L1 · Provenance", "Build gera atestação de procedência (pode ser não assinada)"),
                    ("L2 · Assinada", "Builder hospedado assina a atestação"),
                    ("L3 · Não falsificável", "Builds isolados; agente não forja atestações de outros"),
                    ("L4 · Reprodutível", "Rascunho v1.1: dois builds independentes, bit a bit")],
          "notas": "SLSA (lê-se 'salsa', OpenSSF) é incremental. Mapa de ameaças: dependência maliciosa → materials no L1; build poisoning → ambiente efêmero no L2; falsificação → chave isolada no L3.", "min": 12},
         {"type": "content", "kicker": "NIST SP 800-218", "titulo": "SSDF: quatro grupos de práticas",
          "bullets": [("PO — Prepare the Organization:", "PO.5.1 segregar ambientes"),
                      ("PS — Protect the Software:", "PS.2.1 assinar releases; PS.3.1 integridade"),
                      ("PW — Produce Well-Secured Software:", "PW.4.4 análise estática; PW.4.1 reuso examinado"),
                      ("RV — Respond to Vulnerabilities:", "RV.1.3 monitorar SBOM contra CVEs")],
          "analogia": "SLSA é a nota do teste de resistência do cofre; SSDF é o manual de boas práticas da equipe que cuida do banco. Um mede o artefato, o outro organiza o processo.",
          "notas": "SSDF organiza controles existentes em uma estrutura auditável; é a base do mapeamento de compliance da Aula 8.", "min": 10},
         {"type": "code", "kicker": "Hardening", "titulo": "npm: lockfile, scripts, assinaturas e registry",
          "intro": "Quatro vetores cobertos com .npmrc + npm ci:",
          "code": ["# .npmrc", "package-lock=true", "ignore-scripts=true          # bloqueia postinstall",
                   "audit-level=high", "save-exact=true",
                   "@empresa:registry=https://nexus.empresa.com.br/repository/npm-internal/",
                   "//nexus.empresa.com.br/repository/npm-internal/:always-auth=true", "",
                   "# CI", "git diff --exit-code package-lock.json   # lockfile imutável",
                   "npm ci --ignore-scripts", "npm audit signatures                  # Sigstore",
                   "npm audit --audit-level=high"],
          "notas": "Explique cada linha e o vetor que fecha. Equivalente Python: pip-compile --generate-hashes + pip install --require-hashes. Docker: FROM pinado por @sha256.", "min": 10},
         {"type": "content", "kicker": "Fronteira nova", "titulo": "Riscos emergentes em cadeias de IA",
          "bullets": [("Model poisoning:", "exemplos maliciosos no treinamento → validar origem, hash, drift"),
                      ("Dataset envenenado:", "SBOM para datasets, hash em manifesto"),
                      ("Backdoor em modelo pré-treinado:", "pinning por hash dos pesos, atestação de procedência"),
                      ("Dependência de API de modelo:", "versionamento, SLA, fallback"),
                      ("Licenças em embeddings:", "inventário de modelos no ML-BOM")],
          "notas": "ML-BOM (CycloneDX) inventaria modelos, datasets e parâmetros. Não prometa que SBOM detecta poisoning — ele dá rastreabilidade; a detecção exige validação comportamental.", "min": 8},
         {"type": "lab", "titulo": "Laboratório: hardening de dependências em um pipeline Node.js",
          "passos": [("Mapear:", "fases source/build/distribution e posicionar os vetores"),
                     ("Casos:", "SolarWinds, Codecov, ua-parser-js, XZ — vetor, impacto, controle"),
                     ("Configurar:", ".npmrc com ignore-scripts, save-exact e registry"),
                     ("Verificar:", "npm ci --ignore-scripts, npm audit signatures, npm audit"),
                     ("Mapear:", "cada controle a um nível SLSA e a uma prática SSDF")],
          "entregavel": "Mapa de ameaças, .npmrc e workflow endurecidos, matriz SLSA/SSDF. Runbook: labs/Aula_06_Runbook.md",
          "notas": "Metade discussão, metade técnica. Exija uma evidência por mitigação: hash, assinatura e isolamento respondem a perguntas diferentes.", "min": 70},
         {"type": "recap", "pontos": ["A cadeia tem três fases (origem, build, distribuição) e cada uma tem vetores e controles próprios.",
                                      "Typosquatting, dependency confusion e build poisoning exploram confiança transitiva.",
                                      "SLSA mede garantias de procedência (L1→L4); SSDF organiza práticas (PO, PS, PW, RV).",
                                      "Hardening prático: lockfile com hash, scripts bloqueados, assinaturas verificadas, registry autenticado."],
          "notas": "Pergunta: qual elo tem mais privilégio no seu pipeline? Ponte: Aula 7 prova integridade e origem com Cosign e Sigstore.", "min": 5},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: Integridade e Procedência de Artefatos",
          "notas": "Peça Cosign e um registry local (docker run registry:2) prontos para a próxima aula.", "min": 3},
     ]},

    # ─── AULA 7 ───────────────────────────────────────────────────────────
    {"num": 7, "data": "Sexta, 30/10", "titulo": "Integridade e Procedência de Artefatos",
     "subtitulo": "Como provar que um artefato é autêntico e não foi adulterado",
     "slides": [
         {"type": "content", "kicker": "Problema gerador", "titulo": "docker pull empresa/api:v1.2.3 — e daí?",
          "bullets": [("Uma tag é só bytes:", "não diz quem produziu nem se foi alterado"),
                      ("Entre build e deploy:", "registry, MITM, mirror, replay de versão antiga"),
                      ("Assinatura:", "prova autoria e integridade"),
                      ("Atestação:", "prova propriedades: SBOM, provenance, scan")],
          "analogia": "Um remédio sem lacre e sem bula: pode até ser o original, mas você não tem como saber. Assinatura é o lacre; atestação é a bula com lote, fábrica e composição.",
          "notas": "OWASP A08 — Software and Data Integrity Failures. Pergunte como o consumidor distingue uma imagem legítima de uma substituída no registry. Resposta: digest + verificação.", "min": 8},
         {"type": "comparison", "kicker": "Dois modelos", "titulo": "Chave estática vs. Sigstore keyless",
          "a_title": "Assinatura tradicional (PGP, chave)", "a_items": ["Chave privada de longa duração", "Custódia em HSM/KMS; risco se vazar", "Revogação via CRL/OCSP frágil", "Sem registro público de quem assinou"],
          "b_title": "Sigstore keyless", "b_items": ["Identidade OIDC do pipeline/pessoa", "Certificado efêmero de 10 minutos (Fulcio)", "Chave descartada após assinar", "Transparency log público (Rekor)"],
          "notas": "Keyless não elimina confiança: desloca para o provedor OIDC e para a política de verificação (issuer + subject). Ainda assim, remove o problema da custódia.", "min": 10},
         {"type": "flow", "kicker": "Sigstore", "titulo": "O fluxo keyless em cinco passos",
          "steps": [("Token OIDC", "GitHub Actions, GitLab, Google emitem JWT com identidade"),
                    ("Fulcio", "Emite certificado X.509 de 10 min ligado à identidade"),
                    ("Cosign sign", "Assina o hash do artefato com chave efêmera"),
                    ("Rekor", "Registra assinatura e certificado no log append-only"),
                    ("Verify", "Consumidor valida identidade, digest e entrada no Rekor")],
          "notas": "Percorra o fluxo e destaque: não há chave de longa duração para roubar e toda assinatura fica auditável publicamente.", "min": 12},
         {"type": "code", "kicker": "Cosign", "titulo": "Assinar e verificar uma imagem",
          "intro": "A assinatura é gravada no registry como artefato OCI (tag sha256-<digest>.sig):",
          "code": ["# Keyless (CI ou browser)", "cosign sign --yes ghcr.io/empresa/api@sha256:abc123...", "",
                   "cosign verify ghcr.io/empresa/api@sha256:abc123... \\",
                   "  --certificate-identity-regexp '^https://github.com/empresa/.*' \\",
                   "  --certificate-oidc-issuer https://token.actions.githubusercontent.com", "",
                   "# Chave estática (somente laboratório local)", "cosign generate-key-pair",
                   "cosign sign --key cosign.key localhost:5000/lab/banner:v1",
                   "cosign verify --key cosign.pub localhost:5000/lab/banner:v1"],
          "notas": "Sempre assine o digest, não a tag (tags são mutáveis). A verificação deve restringir issuer e identidade — sem isso, qualquer assinatura válida passa.", "min": 10},
         {"type": "code", "kicker": "Atestações", "titulo": "cosign attest: SBOM e SLSA Provenance", "lang": "bash",
          "intro": "Formato in-toto com envelope DSSE; predicateType diz o que a atestação afirma:",
          "code": ["syft localhost:5000/lab/banner:v1 -o cyclonedx-json > sbom.cdx.json", "",
                   "cosign attest --key cosign.key --type cyclonedx \\",
                   "  --predicate sbom.cdx.json localhost:5000/lab/banner:v1", "",
                   "cosign attest --key cosign.key --type slsaprovenance \\",
                   "  --predicate provenance.json localhost:5000/lab/banner:v1", "",
                   "cosign verify-attestation --key cosign.pub --type slsaprovenance \\",
                   "  localhost:5000/lab/banner:v1 | jq '.payload | @base64d | fromjson | .predicate'", "",
                   "# predicateTypes: cyclonedx · spdx · slsaprovenance · vuln"],
          "notas": "Provenance descreve builder, workflow, commit e dependências. Em produção ela é gerada pelo builder (slsa-github-generator); no lab é demonstrativa.", "min": 10},
         {"type": "code", "kicker": "Verificação como gate", "titulo": "ClusterImagePolicy: só imagem assinada entra", "lang": "yaml",
          "intro": "policy-controller (Sigstore) rejeita no admission o que não bate com a política:",
          "code": ["apiVersion: policy.sigstore.dev/v1beta1", "kind: ClusterImagePolicy",
                   "metadata: { name: require-signed-images }", "spec:",
                   "  images: [{ glob: \"ghcr.io/empresa/**\" }]", "  authorities:",
                   "  - keyless:", "      url: https://fulcio.sigstore.dev", "      identities:",
                   "      - issuer: https://token.actions.githubusercontent.com",
                   "        subjectRegExp: \"^https://github.com/empresa/.*\"",
                   "    ctlog: { url: https://rekor.sigstore.dev }", "    attestations:",
                   "    - { name: must-have-provenance, predicateType: https://slsa.dev/provenance/v1 }", "",
                   "# kubectl label ns producao policy.sigstore.dev/include=true"],
          "notas": "Assinatura só vale se alguém verifica no consumo. Alternativas: Kyverno verifyImages, Connaisseur, Gatekeeper. Também dá para verificar em um job do CI antes do deploy.", "min": 10},
         {"type": "content", "kicker": "Verificação independente", "titulo": "Build reprodutível",
          "bullets": [("Definição:", "mesmo código → artefato bit a bit idêntico"),
                      ("Por quê:", "qualquer um rebuilda e compara com o publicado"),
                      ("Práticas:", "toolchain pinada, SOURCE_DATE_EPOCH, ordenação determinística"),
                      ("Quem faz:", "Debian, Tails, Bitcoin Core; em containers: nixpkgs, apko"),
                      ("SLSA L4:", "detecta comprometimento mesmo com builder 'confiável'")],
          "analogia": "Duas padarias com a mesma receita devem produzir o mesmo pão. Se o seu saiu diferente, alguém mexeu na massa.",
          "notas": "Reprodutibilidade é a verificação que não depende de confiar no builder. Cite também assinatura de modelos de IA (cosign sign model.onnx).", "min": 8},
         {"type": "lab", "titulo": "Laboratório: ciclo completo com Cosign",
          "passos": [("Registry:", "docker run -d -p 5000:5000 registry:2; build e push de banner:v1"),
                     ("Assinar:", "cosign generate-key-pair → sign → verify (chave de teste)"),
                     ("Atestar:", "SBOM CycloneDX (Syft) e provenance demonstrativa"),
                     ("Verificar:", "verify-attestation; conferir sujeito, digest e predicate"),
                     ("Opcional:", "policy-controller no kind e bloqueio de imagem não assinada")],
          "entregavel": "Imagem por digest, assinatura, duas atestações e saídas de verificação. Runbook: labs/Aula_07_Runbook.md",
          "notas": "Chaves de laboratório nunca vão para produção. Diferencie provenance demonstrativa de provenance confiável gerada pelo builder.", "min": 80},
         {"type": "recap", "pontos": ["Assinatura prova autoria e integridade; atestação prova propriedades (SBOM, provenance, scan).",
                                      "Sigstore keyless: OIDC + Fulcio + Rekor eliminam a custódia de chave estática.",
                                      "Assine o digest, verifique com issuer e identidade restritos — no CI e no admission.",
                                      "Build reprodutível é a verificação independente que não depende de confiar no builder."],
          "notas": "Frase-síntese: assinar não basta, o consumidor precisa verificar. Ponte: Aula 8 transforma essas provas em evidência contínua de compliance.", "min": 4},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próxima aula: Compliance Contínuo, Auditoria e Projeto Final",
          "notas": "Lembre a turma de trazer o repositório do projeto final em andamento.", "min": 3},
     ]},

    # ─── AULA 8 ───────────────────────────────────────────────────────────
    {"num": 8, "data": "Terça, 03/11", "titulo": "Compliance Contínuo, Auditoria e Projeto Final",
     "subtitulo": "Como evidenciar controle de ponta a ponta da cadeia em uma auditoria",
     "slides": [
         {"type": "comparison", "kicker": "Problema gerador", "titulo": "Compliance pontual vs. compliance contínuo",
          "a_title": "Pontual", "a_items": ["Campanha anual de coleta de evidências", "Planilhas e capturas de tela", "Alto risco de edição humana", "Desvio detectado meses depois"],
          "b_title": "Contínuo", "b_items": ["Evidência é subproduto do pipeline", "Logs estruturados e assinados na origem", "Baixo risco de falsificação", "Desvio detectado em minutos"],
          "notas": "Pergunte como a organização comprova controles entre auditorias. Compliance contínuo só é viável quando os controles existem como código — é a junção das 7 aulas.", "min": 10},
         {"type": "flow", "kicker": "Arquitetura", "titulo": "Evidência automatizada em quatro camadas",
          "steps": [("Produzir", "OPA decision logs, Cosign bundles, SBOMs, relatórios Trivy"),
                    ("Coletar", "Dependency-Track, OPA bundle service, Loki, Elasticsearch"),
                    ("Imutabilizar", "Object Lock, Rekor, WORM, hash chain"),
                    ("Apresentar", "Grafana, Kibana, relatórios CSV para GRC")],
          "notas": "Cada camada tem componentes típicos. Evidência precisa ser recuperável e verificável de forma independente pelo auditor.", "min": 10},
         {"type": "image", "kicker": "Mapa de evidências", "titulo": "Cada etapa da cadeia deixa uma prova", "image": IMG_SUPPLY,
          "fonte": FONTE_SUPPLY,
          "bullets": [("PR / merge:", "decisão OPA, approver, commit assinado (Gitsign → Rekor)"),
                      ("Build:", "SBOM e provenance anexados via Cosign attest"),
                      ("Registry:", "bundle Sigstore: certificado, assinatura, entrada Rekor"),
                      ("Deploy / runtime:", "decisão do admission controller, eventos Falco, violations")],
          "notas": "Percorra a figura e nomeie a evidência de cada elo. A mesma evidência pode atender vários frameworks ao mesmo tempo.", "min": 10},
         {"type": "content", "kicker": "Frameworks", "titulo": "Um controle, várias exigências",
          "bullets": [("NIST SSDF PS.2.1:", "cosign sign em todo release → atestação no registry"),
                      ("SLSA Build L3:", "builder hospedado + provenance assinada e não falsificável"),
                      ("ISO 27001 A.8.28:", "SAST obrigatório no PR → findings SARIF"),
                      ("PCI DSS 6.2.4:", "Trivy + Dependency-Track → tickets com CVE e correção"),
                      ("LGPD art. 46:", "medidas técnicas documentadas → mapeamento SSDF")],
          "analogia": "Uma única câmera de segurança bem posicionada serve para o seguro, para a polícia e para o condomínio. Evidência boa é reaproveitável.",
          "notas": "Deixe claro: mapear controle a framework não é certificação. É a base probatória que o auditor vai verificar.", "min": 10},
         {"type": "content", "kicker": "Não-repúdio", "titulo": "Trilha de auditoria imutável",
          "bullets": [("Object Lock:", "S3, Azure Blob Immutable, GCS Bucket Lock — retenção de 7 anos"),
                      ("Append-only log:", "Rekor, CloudTrail"),
                      ("Hash chain:", "cada registro referencia o hash do anterior"),
                      ("Assinatura digital:", "Cosign, KMS — ninguém altera depois do fato"),
                      ("Transparência pública:", "Rekor, Certificate Transparency")],
          "notas": "Nem o time interno deve conseguir alterar a evidência. Composição típica: artefato assinado (Rekor) + relatório em bucket com Object Lock + decisões em log com hash chain.", "min": 8},
         {"type": "content", "kicker": "Indicadores", "titulo": "KPIs da postura da cadeia",
          "bullets": [("Releases com SBOM:", "releases_com_sbom / releases_totais → 100%"),
                      ("Imagens assinadas em produção:", "→ 100%"),
                      ("MTTR de CVE crítica:", "data_correção − data_publicação → < 7 dias"),
                      ("Cobertura de admission policies:", "clusters_com_gate / clusters → 100%"),
                      ("Cobertura SLSA L3:", "pipelines_l3 / pipelines → > 80%")],
          "notas": "Todo KPI precisa de numerador, denominador e fonte de dados. Taxa de violação de policy deve ter tendência decrescente, não meta fixa.", "min": 8},
         {"type": "code", "kicker": "Relatório automatizado", "titulo": "Conformidade PS.2.1 em um script",
          "intro": "Um relatório é uma consulta sobre as evidências — com link para o Rekor de cada entrada:",
          "code": ["MES=$(date +%Y-%m)",
                   "gh release list --limit 50 --json tagName,publishedAt \\",
                   "  | jq --arg m \"$MES\" '.[] | select(.publishedAt | startswith($m))' > releases.json", "",
                   "ID='^https://github.com/empresa/.*'",
                   "ISS=https://token.actions.githubusercontent.com", "",
                   "jq -r '.tagName' releases.json | while read TAG; do",
                   "  IMG=\"ghcr.io/empresa/api-pagamentos:$TAG\"",
                   "  SIG=$(cosign verify \"$IMG\" --certificate-identity-regexp \"$ID\" \\",
                   "        --certificate-oidc-issuer \"$ISS\" >/dev/null 2>&1 && echo sim || echo nao)",
                   "  PROV=$(cosign verify-attestation \"$IMG\" --type slsaprovenance \\",
                   "        --certificate-identity-regexp \"$ID\" --certificate-oidc-issuer \"$ISS\" \\",
                   "        >/dev/null 2>&1 && echo sim || echo nao)",
                   "  echo \"$TAG,$SIG,$PROV\" >> conformidade.csv", "done"],
          "notas": "O auditor recebe o CSV e pode validar cada linha de forma independente. O script Python que consome a API do Dependency-Track complementa com CVEs e licenças.", "min": 8},
         {"type": "content", "kicker": "Projeto final", "titulo": "api-pagamentos: requisitos obrigatórios",
          "bullets": [("Policy as Code:", "5+ policies Rego com opa test + 1 policy bloqueante no admission"),
                      ("SBOM:", "CycloneDX e SPDX no pipeline; ingestão no Dependency-Track; allowlist de licenças"),
                      ("Cadeia:", "Cosign keyless, SLSA Provenance anexada, gate que rejeita imagem sem assinatura"),
                      ("Auditoria:", "relatório mensal automatizado + dashboard com 4 KPIs"),
                      ("Documentação:", "README com mapeamento NIST SSDF, SLSA e ISO 27001")],
          "notas": "Mostre a estrutura do repositório e os workflows release.yml e verify.yml. Rubrica: insuficiente (0–39) a excelente (90–100); pipeline sem intervenção manual e apresentação valem 15 pontos cada.", "min": 12},
         {"type": "flow", "kicker": "Banca", "titulo": "Apresentação em 15 minutos",
          "steps": [("Contexto · 2 min", "A aplicação e os riscos que enfrenta"),
                    ("Demo · 3 min", "Pipeline do commit ao deploy"),
                    ("Ataque · 3 min", "Imagem não assinada → bloqueio no gate"),
                    ("Evidência · 2 min", "Prova pedida pela banca: assinatura, SBOM ou provenance"),
                    ("Framework + Q&A · 5 min", "Mapeamento escolhido pela banca e perguntas")],
          "notas": "O cenário adversarial é obrigatório: capture a saída do admission webhook negando como evidência. A banca pode pedir qualquer evidência — tudo precisa estar recuperável.", "min": 8},
         {"type": "lab", "titulo": "Laboratório: pipeline final integrador",
          "passos": [("Integrar:", "policies Rego + SBOM (Syft) + scan (Trivy) no release.yml"),
                     ("Assinar e atestar:", "cosign sign + attest (SBOM, vuln) + slsa-github-generator"),
                     ("Gate:", "ClusterImagePolicy exigindo assinatura, SBOM e provenance"),
                     ("Adversarial:", "kubectl run intruso --image=...:fake → denied"),
                     ("Evidenciar:", "gera_relatorio.sh, consulta_dt.py e tabela de frameworks no README")],
          "entregavel": "Repositório funcional, evidências do bloqueio, CSV de conformidade e README mapeado. Runbook: labs/Aula_08_Runbook.md",
          "notas": "Distribua responsabilidades por camada nos grupos. Cada grupo deve provar uma decisão com log, digest ou atestação antes de encerrar.", "min": 70},
         {"type": "recap", "pontos": ["Compliance contínuo: evidência como subproduto do pipeline, não como campanha anual.",
                                      "Quatro camadas: produzir, coletar, imutabilizar, apresentar — com não-repúdio em cada uma.",
                                      "A mesma evidência atende NIST SSDF, SLSA, ISO 27001, PCI DSS e LGPD.",
                                      "Projeto final integra PaC + SBOM + Cosign + relatórios; a banca verifica evidências, não narrativas."],
          "notas": "Recapitule a progressão do módulo: política → inventário → cadeia → compliance. Combine horário e ordem das apresentações.", "min": 6},
         {"type": "closing", "msg": "Obrigado!", "sub": "Próximo encontro: concluir atividades e preencher rubricas",
          "notas": "Reforce prazos de entrega do projeto e das atividades do LMS.", "min": 3},
     ]},

    # ─── AULA 9 ───────────────────────────────────────────────────────────
    {"num": 9, "data": "Sexta, 06/11", "titulo": "Concluir Atividades e Preencher Rubricas",
     "subtitulo": "Encerramento do módulo, apresentações finais e avaliação",
     "slides": [
         {"type": "content", "kicker": "Encerramento", "titulo": "Objetivos do último encontro",
          "bullets": [("Apresentações:", "projeto final em 15 minutos por grupo, com banca"),
                      ("Concluir atividades:", "exercícios pendentes do LMS"),
                      ("Revisão geral:", "amarrar política, inventário, cadeia e compliance"),
                      ("Preencher rubricas:", "avaliação do aprendizado e feedback do módulo")],
          "analogia": "É a auditoria final do módulo: em vez de contar o que aprendeu, você mostra a evidência — o pipeline rodando e o gate bloqueando.",
          "notas": "Defina a ordem das apresentações e o tempo. Lembre que a banca pode pedir qualquer evidência do pipeline.", "min": 8},
         {"type": "flow", "kicker": "Revisão", "titulo": "A progressão do módulo",
          "steps": [("Política", "OPA, Rego, Conftest, Gatekeeper — Aulas 1 a 3"),
                    ("Inventário", "SBOM, SPDX/CycloneDX, Dependency-Track — Aulas 4 e 5"),
                    ("Cadeia", "Ameaças, SLSA, SSDF, Cosign, Sigstore — Aulas 6 e 7"),
                    ("Compliance", "Evidência contínua, auditoria, KPIs — Aula 8")],
          "notas": "Peça que cada grupo diga, em uma frase, qual controle de cada camada implementou no projeto.", "min": 10},
         {"type": "content", "kicker": "Revisão", "titulo": "Perguntas que vocês agora sabem responder",
          "bullets": [("Como transformar uma regra em código?", "Rego + OPA + gate no pipeline"),
                      ("Como impor no cluster?", "Gatekeeper: Template, Constraint, dryrun → deny"),
                      ("O que tem dentro do software?", "SBOM com PURL, correlacionado no Dependency-Track"),
                      ("Como o atacante entra sem tocar no código?", "Dependências, build, distribuição — SLSA e SSDF"),
                      ("Como provar que é autêntico?", "Cosign, Sigstore, atestações; e verificação no consumo")],
          "notas": "Use como quiz rápido antes das apresentações. Ótimo termômetro para as rubricas.", "min": 10},
         {"type": "content", "kicker": "Avaliação", "titulo": "Como preencher as rubricas",
          "bullets": [("Autoavaliação:", "reflita sobre o que consegue fazer hoje que não conseguia"),
                      ("Critérios:", "cada rubrica avalia uma competência específica do módulo"),
                      ("Honestidade:", "avalie com sinceridade para evoluir de verdade"),
                      ("Feedback:", "o que funcionou e o que mudar no módulo")],
          "analogia": "A rubrica é o SBOM do seu aprendizado: lista o que você tem, em que versão, e onde ainda há 'CVE' para corrigir.",
          "notas": "Reserve tempo real para o preenchimento. Fique disponível para dúvidas individuais.", "min": 15},
         {"type": "lab", "titulo": "Apresentações do projeto final", "kicker": "Banca",
          "passos": [("Contexto:", "aplicação e riscos — 2 min"),
                     ("Demo:", "pipeline do commit ao deploy — 3 min"),
                     ("Ataque:", "imagem não assinada bloqueada — 3 min"),
                     ("Evidência + framework:", "pedidos pela banca — 5 min"),
                     ("Perguntas:", "2 min")],
          "entregavel": "Repositório, README com mapeamento de frameworks, evidências e apresentação. Rubrica: 0–100 pontos.",
          "notas": "Mantenha o tempo com rigor. Anote evidências apresentadas para fundamentar a nota.", "min": 105},
         {"type": "recap", "pontos": ["Política → inventário → cadeia → compliance: uma camada de governança sobre o pipeline DevSecOps.",
                                      "Controles viram código, código vira gate, gate vira evidência.",
                                      "As rubricas registram o aprendizado e orientam os próximos passos.",
                                      "Parabéns! Você concluiu o módulo CY-SE-013 — Supply Chain Security! 🎉"],
          "notas": "Agradeça a participação e indique caminhos de aprofundamento: OpenSSF, SLSA, Sigstore, CycloneDX.", "min": 9},
         {"type": "closing", "msg": "Parabéns e obrigado!", "sub": "Módulo concluído — continuem verificando a cadeia.",
          "notas": "Encerramento. Deixe contatos e materiais de referência.", "min": 3},
     ]},
]


# =============================================================================
#  MONTAGEM
# =============================================================================
def build_aula(aula):
    num, titulo = aula["num"], aula["titulo"]
    prs = new_prs()
    idx = 1
    sl = build_capa(prs, num, titulo, aula["subtitulo"], aula["data"])
    set_notes(sl, f"Abertura da Aula {num:02d}: {titulo}. Apresente o tema, a pergunta geradora e situe a aula na progressão do módulo. Confirme que os laboratórios da aula anterior foram concluídos.", 3)
    idx += 1
    sl = build_cronograma(prs, num, titulo, idx)
    set_notes(sl, "Mostre onde estamos no cronograma (card escuro). Lembre os próximos encontros e, na Aula 7, avise que não há aula em 02/11 (Finados). O último encontro, 06/11, é dedicado às apresentações e rubricas.", 2)
    idx += 1
    sl = build_combinado(prs, titulo, idx)
    set_notes(sl, "Relembre os combinados da turma rapidamente: câmeras ligadas, pontualidade, participação e mão levantada para dúvidas. Informe o horário do intervalo de 15 minutos.", 2)
    for sd in aula["slides"]:
        idx += 1
        t = sd["type"]
        if t == "content":
            sl = build_content(prs, titulo, idx, sd["titulo"], sd["bullets"], sd.get("analogia"), sd.get("kicker"))
        elif t == "image":
            sl = build_image(prs, titulo, idx, sd["titulo"], sd["image"], sd.get("bullets"), sd.get("caption"), sd.get("kicker"), sd.get("fonte"))
        elif t == "flow":
            sl = build_flow(prs, titulo, idx, sd["titulo"], sd["steps"], sd.get("kicker"), sd.get("analogia"))
        elif t == "code":
            sl = build_code(prs, titulo, idx, sd["titulo"], sd.get("intro"), sd["code"], sd.get("kicker"), sd.get("lang", "bash"))
        elif t == "comparison":
            sl = build_comparison(prs, titulo, idx, sd["titulo"], sd["a_title"], sd["a_items"], sd["b_title"], sd["b_items"], sd.get("kicker"))
        elif t == "lab":
            sl = build_lab(prs, titulo, idx, sd["titulo"], sd["passos"], sd["entregavel"], sd.get("kicker", "Mão na massa"))
        elif t == "recap":
            sl = build_recap(prs, titulo, idx, sd["pontos"])
        elif t == "closing":
            sl = build_closing(prs, sd["msg"], sd["sub"])
        set_notes(sl, sd["notas"], sd["min"])
    safe = "".join(c for c in unicodedata.normalize("NFKD", titulo) if ord(c) < 128)
    safe = safe.replace(":", "").replace(",", "").replace(" ", "_").replace("/", "-")
    out = os.path.join(BASE, f"Aula_{num:02d}_{safe}.pptx")
    prs.save(out)
    total = 3 + 2 + 2 + sum(s["min"] for s in aula["slides"])
    print(f"Saved: {os.path.basename(out)}  ({len(prs.slides._sldIdLst)} slides, {total} min)")
    return out


if __name__ == "__main__":
    for aula in AULAS:
        build_aula(aula)
