"""
add_review_slides.py
Inserts a styled "Revisão da Aula Anterior" slide at index 2 (0-based)
in each course PPTX: Capa=0, Quebra-Gelo=1, Revisão=2 (NEW), content...
For Aula 01, adds "Visão Geral do Curso" instead.

Run once: python slides/add_review_slides.py
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

C_BG    = RGBColor(0x18, 0x1C, 0x23)
C_WHITE = RGBColor(0xFF, 0xFD, 0xF8)
C_BLUE  = RGBColor(0x23, 0x71, 0xFF)
C_GREEN = RGBColor(0xA6, 0xF7, 0x50)
C_DGRAY = RGBColor(0x2A, 0x2F, 0x3A)

SW, SH = 9144000, 5143500
SLIDES_DIR = r"c:\UserData\z004stze\Documents\Private\mine\ada\orquestracao-k8s\slides"


# ── Drawing helpers ───────────────────────────────────────────────────────────

def _shape(slide, kind, x, y, w, h, fill=None, line=None):
    s = slide.shapes.add_shape(kind, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = fill
    else:    s.fill.background()
    if line: s.line.color.rgb = line
    else:    s.line.fill.background()
    return s

def rect(slide, x, y, w, h, fill=None, line=None):
    return _shape(slide, 1, x, y, w, h, fill, line)

def rrect(slide, x, y, w, h, fill=None, line=None, radius=5):
    s = _shape(slide, 5, x, y, w, h, fill, line)
    pg = s.element.find('.//' + qn('a:prstGeom'))
    if pg is not None:
        av = pg.find(qn('a:avLst'))
        if av:
            for gd in av.findall(qn('a:gd')):
                if gd.get('name') == 'adj':
                    gd.set('fmla', f'val {radius * 1000}')
    return s

def oval(slide, x, y, w, h, fill=None, line=None):
    return _shape(slide, 9, x, y, w, h, fill, line)

def tb(slide, x, y, w, h, text, font="Ubuntu", sz=12,
       bold=False, color=None, align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = font; run.font.size = Pt(sz); run.font.bold = bold
    if color: run.font.color.rgb = color
    return box


# ── Review content per aula ───────────────────────────────────────────────────
# Each item is (label, value) — label shown in green, value in white.

REVIEWS = {
    "Aula_01": {
        "badge":    "VISÃO GERAL DO CURSO",
        "title":    "O que vamos construir juntos?",
        "subtitle": "8 aulas · do Minikube local ao Cluster Multi-Master em alta disponibilidade",
        "items": [
            ("Aula 1",  "Conhecendo o Kubernetes — arquitetura e fundamentos"),
            ("Aula 2",  "Mão na Massa — Minikube e kubectl no seu computador"),
            ("Aula 3",  "Pods, Namespaces e Controladores — objetos essenciais"),
            ("Aula 4",  "Primeira Aplicação no K8S — deploy completo do zero"),
            ("Aula 5",  "Estratégias de Deploy e Limitação de Recursos"),
            ("Aula 6",  "Services, Ingress e Volumes — redes e persistência"),
            ("Aula 7",  "Componentes e Troubleshooting — diagnóstico e RBAC"),
            ("Aula 8",  "Cluster Multi-Master — Alta Disponibilidade com etcd HA"),
        ],
    },
    "Aula_02": {
        "badge":    "REVISÃO — AULA 1",
        "title":    "Conhecendo o Kubernetes",
        "subtitle": "Fundamentos que precisam estar sólidos antes de avançar",
        "items": [
            ("Origem",         "Google Borg (2003) → K8s open-source (2014) → doado à CNCF"),
            ("Problema",       "Orquestração: scheduling, self-healing, scaling, service discovery"),
            ("Control Plane",  "API Server · etcd · Scheduler · Controller Manager"),
            ("Worker Node",    "kubelet · kube-proxy · Container Runtime (containerd)"),
            ("Objetos base",   "Pod · ReplicaSet · Deployment · Service · Namespace"),
            ("CNCF",           "K8s é o projeto âncora do ecossistema Cloud Native Foundation"),
        ],
    },
    "Aula_03": {
        "badge":    "REVISÃO — AULA 2",
        "title":    "Mão na Massa: Minikube e kubectl",
        "subtitle": "Seu lab local está configurado — revisando ferramentas e comandos essenciais",
        "items": [
            ("Minikube",       "Cluster K8s local single-node para desenvolvimento e estudo"),
            ("kubectl",        "CLI oficial: get · describe · apply · delete · logs · exec"),
            ("Contextos",      "kubeconfig gerencia múltiplos clusters (kubectl config use-context)"),
            ("Dashboard",      "minikube dashboard → UI visual do cluster no browser"),
            ("Namespaces",     "default · kube-system · kube-public · kube-node-lease"),
            ("Boas práticas",  "kubectl apply (declarativo, idempotente) é preferível ao create"),
        ],
    },
    "Aula_04": {
        "badge":    "REVISÃO — AULA 3",
        "title":    "Pods, Namespaces e Controladores",
        "subtitle": "A hierarquia de objetos do K8s e como eles se organizam",
        "items": [
            ("Pod",            "Menor unidade — 1+ containers com IP e volumes compartilhados"),
            ("Ciclo de vida",  "Pending → Running → Succeeded / Failed / Unknown"),
            ("Namespace",      "Isolamento lógico — uma caixa por equipe ou ambiente"),
            ("ReplicaSet",     "Garante N réplicas rodando — recria pods automaticamente"),
            ("Deployment",     "Gerencia ReplicaSets — habilita rolling update e rollback"),
            ("Labels/Select.", "Como objetos se encontram: spec.selector.matchLabels define o vínculo"),
        ],
    },
    "Aula_05": {
        "badge":    "REVISÃO — AULA 4",
        "title":    "Primeira Aplicação no K8S",
        "subtitle": "Do Dockerfile ao cluster — o fluxo completo de um deploy",
        "items": [
            ("YAML",             "apiVersion · kind · metadata · spec — os 4 campos obrigatórios"),
            ("apply vs create",  "apply = declarativo (idempotente) · create = imperativo (1x)"),
            ("Deployment YAML",  "spec.replicas · spec.selector · spec.template.spec.containers"),
            ("Service NodePort", "Expõe o pod externamente via porta do nó (30000–32767)"),
            ("ConfigMap",        "Variáveis de config externalizadas — sem rebuild de imagem"),
            ("Secret",           "Dados sensíveis em base64 — nunca versionar no repositório"),
        ],
    },
    "Aula_06": {
        "badge":    "REVISÃO — AULA 5",
        "title":    "Estratégias de Deploy e Limitação de Recursos",
        "subtitle": "Deployar com segurança, escalar com inteligência",
        "items": [
            ("Rolling Update",  "Substituição gradual de pods — zero downtime (padrão do K8s)"),
            ("Blue/Green",      "Dois ambientes paralelos — switch instantâneo pelo Service"),
            ("Canary",          "Tráfego dividido — ex: 90% v1 / 10% v2 para teste progressivo"),
            ("Probes",          "Liveness (vivo?) · Readiness (pronto p/ tráfego?) · Startup"),
            ("Requests/Limits", "Request = reserva mínima · Limit = teto máximo (CPU/RAM)"),
            ("HPA",             "HorizontalPodAutoscaler — escala réplicas automaticamente por métricas"),
        ],
    },
    "Aula_07": {
        "badge":    "REVISÃO — AULA 6",
        "title":    "Services, Ingress e Volumes",
        "subtitle": "Redes e persistência — conectando e armazenando dados no cluster",
        "items": [
            ("ClusterIP",    "IP estável interno — pods se comunicam por nome do Service"),
            ("NodePort / LB","NodePort: porta no nó · LoadBalancer: IP externo via cloud provider"),
            ("CoreDNS",      "DNS interno: <svc>.<namespace>.svc.cluster.local"),
            ("Ingress",      "Roteamento HTTP/HTTPS por domínio/path para múltiplos Services"),
            ("PV / PVC",     "PersistentVolume (disco real) ← PersistentVolumeClaim (pedido)"),
            ("StorageClass", "Provisionamento dinâmico — PV criado automaticamente pelo cluster"),
        ],
    },
    "Aula_08": {
        "badge":    "REVISÃO — AULA 7",
        "title":    "Componentes e Troubleshooting",
        "subtitle": "As entranhas do K8s e como diagnosticar falhas com método",
        "items": [
            ("RBAC",             "Role · ClusterRole · RoleBinding · ServiceAccount — controle de acesso"),
            ("CrashLoopBackOff", "App termina logo após iniciar — verifique logs e variáveis de ambiente"),
            ("OOMKilled",        "Container estourou o memory limit — aumente limits ou otimize uso"),
            ("ImagePullBackOff", "Imagem não encontrada ou credenciais inválidas no registry"),
            ("kubectl debug",    "Injeta container sidecar no pod em execução para diagnóstico ao vivo"),
            ("Checklist",        "logs → describe → events → exec → debug → kubectl top"),
        ],
    },
}

FILES = [
    ("Aula_01", "Aula_01_Conhecendo_o_Kubernetes.pptx"),
    ("Aula_02", "Aula_02_Mao_na_Massa_Minikube_e_kubectl.pptx"),
    ("Aula_03", "Aula_03_Pods_Namespaces_e_Controladores.pptx"),
    ("Aula_04", "Aula_04_Primeira_Aplicacao_no_K8S.pptx"),
    ("Aula_05", "Aula_05_Estrategias_de_Deploy_e_Limitacao_de_Recursos.pptx"),
    ("Aula_06", "Aula_06_Services_Ingress_e_Volumes.pptx"),
    ("Aula_07", "Aula_07_Componentes_e_Troubleshooting.pptx"),
    ("Aula_08", "Aula_08_Cluster_Multi_Master.pptx"),
]


# ── Slide builder ─────────────────────────────────────────────────────────────

def build_review_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Background + left accent bar
    rect(slide, 0, 0, SW, SH, fill=C_BG)
    rect(slide, 0, 0, 55000, SH, fill=C_GREEN)

    # Header region
    HEADER_H = 1050000
    rect(slide, 55000, 0, SW - 55000, HEADER_H, fill=C_DGRAY)

    # Badge pill
    BADGE_W, BADGE_H = 1850000, 165000
    BADGE_X, BADGE_Y = 130000, 130000
    rrect(slide, BADGE_X, BADGE_Y, BADGE_W, BADGE_H, fill=C_BLUE, radius=50)
    tb(slide, BADGE_X + 65000, BADGE_Y + 32000, BADGE_W - 130000, BADGE_H - 64000,
       data["badge"], sz=8, bold=True, color=C_WHITE)

    # Title
    tb(slide, BADGE_X, BADGE_Y + BADGE_H + 55000,
       SW - BADGE_X - 380000, 390000,
       data["title"], sz=26, bold=True, color=C_WHITE)

    # Subtitle
    tb(slide, BADGE_X, BADGE_Y + BADGE_H + 55000 + 360000,
       SW - BADGE_X - 380000, 230000,
       data["subtitle"], sz=10, color=C_GREEN)

    # Decorative dots (top-right, matching Ada style)
    for dx, dy, dr in [(SW - 320000, 180000, 55000),
                       (SW - 190000, 340000, 35000),
                       (SW - 430000, 130000, 28000)]:
        oval(slide, dx - dr, dy - dr, 2 * dr, 2 * dr, fill=C_GREEN)

    # Content area — 2-column card grid
    CY   = HEADER_H + 85000
    CH   = SH - CY - 70000
    MGN  = 120000
    CGAP = 55000
    RGAP = 45000
    TOTW = SW - 55000 - 2 * MGN

    items  = data["items"]
    NCOLS  = 2
    NROWS  = -(-len(items) // NCOLS)          # ceiling division
    COL_W  = (TOTW - (NCOLS - 1) * CGAP) // NCOLS
    COL1_X = 55000 + MGN
    COL2_X = COL1_X + COL_W + CGAP
    ROW_H  = (CH - (NROWS - 1) * RGAP) // NROWS

    for idx, (label, value) in enumerate(items):
        col = idx % NCOLS
        row = idx // NCOLS
        cx = COL1_X if col == 0 else COL2_X
        cy = CY + row * (ROW_H + RGAP)

        # Card background
        rrect(slide, cx, cy, COL_W, ROW_H, fill=C_DGRAY, radius=3)

        # Green left accent bar inside card
        rrect(slide, cx + 25000, cy + 30000, 18000, ROW_H - 60000,
              fill=C_GREEN, radius=10)

        TEXT_X = cx + 25000 + 18000 + 35000
        TEXT_W = COL_W - 25000 - 18000 - 35000 - 30000

        LABEL_H = 175000
        tb(slide, TEXT_X, cy + 35000, TEXT_W, LABEL_H,
           label, sz=8, bold=True, color=C_GREEN)
        tb(slide, TEXT_X, cy + LABEL_H + 15000, TEXT_W, ROW_H - LABEL_H - 50000,
           value, sz=9, color=C_WHITE, wrap=True)

    # Bottom blue bar
    rect(slide, 55000, SH - 55000, SW - 55000, 55000, fill=C_BLUE)

    return slide


def insert_at(prs, target_index):
    """Move the last slide (just appended) to target_index."""
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    last = items[-1]
    sldIdLst.remove(last)
    sldIdLst.insert(target_index, last)


def _slide_badge_text(slide):
    """Return all text from a slide concatenated, for duplicate detection."""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    return " ".join(texts)

def _remove_duplicate_review(prs, badge):
    """If index 2 AND index 3 both contain the badge, remove index 3."""
    slides_list = prs.slides._sldIdLst
    if len(prs.slides) <= 3:
        return
    t2 = _slide_badge_text(prs.slides[2])
    t3 = _slide_badge_text(prs.slides[3])
    if badge in t2 and badge in t3:
        # remove index 3
        elem = list(slides_list)[3]
        slides_list.remove(elem)
        print(f"    Removed duplicate review slide at index 3")

def process(aula_key, filename):
    path = os.path.join(SLIDES_DIR, filename)
    if not os.path.exists(path):
        print(f"  NOT FOUND: {filename}")
        return
    data = REVIEWS.get(aula_key)
    if not data:
        print(f"  NO DATA: {aula_key}")
        return

    prs = Presentation(path)
    n = len(prs.slides)
    badge = data["badge"]

    # Idempotency: fix duplicates first, then check if already present
    _remove_duplicate_review(prs, badge)

    if n > 2 and badge in _slide_badge_text(prs.slides[2]):
        print(f"  SKIP {filename}  (review slide already at index 2)")
        prs.save(path)  # save dedup fix if any
        return

    insert_idx = min(2, n)
    print(f"  {filename}  ({n} slides) -> inserting at index {insert_idx}")

    build_review_slide(prs, data)
    insert_at(prs, insert_idx)
    prs.save(path)
    print(f"    Saved OK  (now {len(prs.slides)} slides)")


if __name__ == "__main__":
    print("=== Adding Review / Overview slides ===\n")
    for aula_key, filename in FILES:
        process(aula_key, filename)
    print("\nAll done!")
